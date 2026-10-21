"""Rebuild a frozen 2025 operations plan. No email, payroll or live release.

Team capacity and scheduling conventions are original task inputs, not law.
"""
import argparse,collections,csv,datetime as dt,hashlib,json,shutil,subprocess,sys,tempfile
from pathlib import Path
from docx import Document
from docx.shared import Cm,Pt
from docx.oxml import OxmlElement
from openpyxl import Workbook,load_workbook
from openpyxl.styles import Font,PatternFill,Alignment
from openpyxl.utils import get_column_letter

INPUTS=['S01_2013旧办法.docx','S02_2024修订.docx','S03_2025年度安排.docx','S04_官方日历摘录.json','S05_未审预排v0.xlsx','S06_运营委托与依赖合同.docx','S07_岗位表.csv','S08_请假培训.csv','S09_公司停服.csv','S10_客户里程碑.csv','S11_版本冻结登记.json','S12_公农历核对.docx']
OUTPUTS=['operations_plan.xlsx','client_messages.docx','governance_note.docx','workflow.py']
DAG={'sources':[],'calendar':['sources'],'service':['calendar'],'availability':['service'],'meetings':['service'],'allocation':['availability','meetings'],'documents':['allocation','calendar'],'preflight':['documents']}

def canonical(x):return json.dumps(x,ensure_ascii=False,sort_keys=True,separators=(',',':')).encode()
def digest(x):return hashlib.sha256(x).hexdigest()
def sha(p):return digest(Path(p).read_bytes())
def readcsv(p):
    with p.open(encoding='utf-8-sig',newline='') as f:return list(csv.DictReader(f))
def writejson(p,x):p.write_text(json.dumps(x,ensure_ascii=False,indent=2)+'\n')
def days(start,end):
    a,b=dt.date.fromisoformat(start),dt.date.fromisoformat(end)
    if a>b:raise ValueError('reversed date interval')
    return [(a+dt.timedelta(days=i)).isoformat() for i in range((b-a).days+1)]
def snapshots(inp):return {n:sha(inp/n) for n in INPUTS}

def frozen_sources(inp):
    facts=json.loads((inp/'S04_官方日历摘录.json').read_text())
    binding=json.loads((inp/'S11_版本冻结登记.json').read_text())
    protected={n for n in INPUTS if n[:3] in {'S01','S02','S03','S04','S06','S12'}}
    if facts['year']!=2025 or binding['status']!='reviewed_for_planning':
        raise ValueError('unapproved source version')
    if set(binding['source_hashes'])!=protected:
        raise ValueError('source freeze coverage')
    for n,h in binding['source_hashes'].items():
        if sha(inp/n)!=h:raise ValueError('source freeze drift: '+n)
    return facts,binding

def national(facts,legacy):
    year=facts['year'];dates=days(f'{year}-01-01',f'{year}-12-31')
    statutory=set(facts['statutory_dates']);makeup=set(facts['makeup_workdays'])
    holidays={d:part['name'] for part in facts['rest_windows'] for d in days(part['start'],part['end'])}
    if not statutory<=set(holidays) or makeup&set(holidays):raise ValueError('inconsistent holiday definitions')
    if set(legacy)!=set(dates):raise ValueError('legacy calendar incomplete')
    rows=[]
    for date in dates:
        weekday=dt.date.fromisoformat(date).weekday()
        kind='STATUTORY_HOLIDAY' if date in statutory else 'ADJUSTED_REST' if date in holidays else 'MAKEUP_WORKDAY' if date in makeup else 'WEEKEND_REST' if weekday>=5 else 'REGULAR_WORKDAY'
        work=int(kind in ['MAKEUP_WORKDAY','REGULAR_WORKDAY'])
        rows.append(dict(date=date,weekday=weekday+1,kind=kind,statutory=int(date in statutory),national_work=work,legacy_work=legacy[date],delta=work-legacy[date],source='S02第二条 / S03对应节日段'))
    return rows

def service(calendar,closures):
    shut={r['date']:r['reason'] for r in closures if r['status']=='approved'}
    return [dict(date=r['date'],company_open=int(r['national_work'] and r['date'] not in shut),national_work=r['national_work'],local_reason=shut.get(r['date'],'')) for r in calendar]

def availability(company,staff,leave):
    ids={r['person'] for r in staff}
    if len(ids)!=len(staff) or any(r['role'] not in ['analyst','reviewer'] for r in staff):raise ValueError('invalid roster')
    if any(r['person'] not in ids for r in leave):raise ValueError('unknown staff in leave register')
    blocked=collections.defaultdict(list)
    for r in leave:
        if r['status']=='approved':blocked[r['date'],r['person']].append(r['reason'])
        elif r['status']!='pending':raise ValueError('unknown approval state')
    return [dict(date=c['date'],person=s['person'],role=s['role'],available=int(c['company_open'] and not blocked[c['date'],s['person']]),reason=';'.join(blocked[c['date'],s['person']]) or ('公司未营业' if not c['company_open'] else '可用')) for c in company for s in staff]

def resolve_meetings(company,requests):
    open_days=[r['date'] for r in company if r['company_open']];inside={r['date'] for r in company};out=[];seen=set()
    for r in requests:
        if r['id'] in seen:raise ValueError('duplicate project ID')
        seen.add(r['id']);dt.date.fromisoformat(r['requested'])
        if r['priority'] not in ['P1','P2'] or r['mode'] not in ['fixed','flexible']:raise ValueError('invalid request type')
        meeting=None;status='UNALLOCATED'
        if r['requested'] not in inside:status='OUT_OF_SCOPE'
        elif r['mode']=='fixed':
            if r['requested'] not in open_days:status='BLOCKED_DATE'
            else:meeting=r['requested']
        else:
            meeting=next((d for d in open_days if d>=r['requested']),None)
            if meeting is None:status='BLOCKED_DATE'
        out.append(dict(r,meeting=meeting,status=status,client_confirmation_required=r['mode']=='flexible'))
    return out

def allocate(people,meetings):
    pool=collections.defaultdict(list)
    for r in people:
        if r['available']:pool[r['date'],r['role']].append(r['person'])
    dates=sorted({r['date'] for r in people});used=set();jobs=[];plans=[]
    for r in sorted(meetings,key=lambda r:(r['priority'],r['meeting'] or '9999',r['id'])):
        p=dict(r,prep_date=None,preparer=None,review_date=None,reviewer=None)
        if r['status']!='UNALLOCATED':plans.append(p);continue
        found=None
        for review_day in reversed([d for d in dates if d<r['meeting']]):
            reviewers=[s for s in sorted(pool[review_day,'reviewer']) if (review_day,s) not in used]
            if not reviewers:continue
            for prep_day in reversed([d for d in dates if d<review_day]):
                analysts=[s for s in sorted(pool[prep_day,'analyst']) if (prep_day,s) not in used]
                if analysts:found=prep_day,analysts[0],review_day,reviewers[0];break
            if found:break
        if found is None:p['status']='BLOCKED_CAPACITY'
        else:
            prep,person,review,checker=found;used.update([(prep,person),(review,checker)])
            p.update(prep_date=prep,preparer=person,review_date=review,reviewer=checker,status='PROPOSED_CLIENT_CONFIRMATION' if r['mode']=='flexible' else 'PLANNED_FIXED')
            jobs.extend([dict(project=r['id'],phase='prepare',date=prep,person=person,role='analyst'),dict(project=r['id'],phase='review',date=review,person=checker,role='reviewer')])
        plans.append(p)
    return dict(projects=sorted(plans,key=lambda p:p['id']),jobs=sorted(jobs,key=lambda r:(r['date'],r['person'],r['project'])))

def validate_allocation(plan,people):
    allowed={(r['date'],r['person'],r['role']) for r in people if r['available']};seen=set()
    known={p['id'] for p in plan['projects']}
    if len(known)!=len(plan['projects']):raise ValueError('duplicate planned project')
    for job in plan['jobs']:
        if job['project'] not in known:raise ValueError('unknown project job')
        if (job['date'],job['person'],job['role']) not in allowed:raise ValueError('assignment on unavailable slot')
        key=job['date'],job['person']
        if key in seen:raise ValueError('double booking')
        seen.add(key)
    for p in plan['projects']:
        if p['status'] in ['PLANNED_FIXED','PROPOSED_CLIENT_CONFIRMATION']:
            if not(p['prep_date']<p['review_date']<p['meeting']) or p['preparer']==p['reviewer']:raise ValueError('invalid preparation/review dependency')
            expected=[dict(project=p['id'],phase='prepare',date=p['prep_date'],person=p['preparer'],role='analyst'),dict(project=p['id'],phase='review',date=p['review_date'],person=p['reviewer'],role='reviewer')]
            actual=[j for j in plan['jobs'] if j['project']==p['id']]
            if sorted(actual,key=lambda j:j['phase'])!=sorted(expected,key=lambda j:j['phase']):raise ValueError('project/job records disagree')
        elif any(p[k] is not None for k in ['prep_date','review_date','preparer','reviewer']):raise ValueError('blocked project has booked resources')
        elif any(j['project']==p['id'] for j in plan['jobs']):raise ValueError('blocked project has job rows')

def word(path,title,sections):
    doc=Document();doc.styles['Normal'].font.name='Noto Sans CJK SC';doc.styles['Normal'].font.size=Pt(10)
    for sec in doc.sections:sec.top_margin=sec.bottom_margin=Cm(1.7)
    doc.add_heading(title,0)
    for head,paragraphs,rows in sections:
        doc.add_heading(head,1)
        for text in paragraphs:doc.add_paragraph(text)
        if rows:
            table=doc.add_table(rows=1,cols=len(rows[0]));table.style='Table Grid';table.autofit=False
            widths=[1.4,5.4,9.7] if len(rows[0])==3 else [16.5/len(rows[0])]*len(rows[0])
            for col,width in zip(table.columns,widths):col.width=Cm(width)
            for c,v in zip(table.rows[0].cells,rows[0]):c.text=str(v)
            for row in rows[1:]:
                tr=table.add_row();tr._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
                for c,v in zip(tr.cells,row):c.text=str(v or '')
            table.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
            for row in table.rows:
                for cell,width in zip(row.cells,widths):
                    cell.width=Cm(width)
                    for para in cell.paragraphs:
                        for run in para.runs:run.font.size=Pt(8.5)
    doc.sections[0].footer.paragraphs[0].text='原创训练情境；待内部复核，未发送通知或执行真实排班。'
    doc.save(path)

def workbook(path,cal,company,people,plan,facts):
    w=Workbook();w.remove(w.active)
    def sheet(name,headers,rows):
        s=w.create_sheet(name);s.append(headers)
        for row in rows:s.append(row)
        s.freeze_panes='A2';s.auto_filter.ref=s.dimensions
        for c in s[1]:c.font=Font(bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='173C45')
        for row in s:
            for c in row:c.alignment=Alignment(vertical='top',wrap_text=True)
        for i in range(1,len(headers)+1):s.column_dimensions[get_column_letter(i)].width=20
        s.row_dimensions[1].height=32
        s.print_title_rows='1:1';s.page_setup.orientation='landscape';s.page_setup.paperSize=s.PAPERSIZE_A3;s.sheet_properties.pageSetUpPr.fitToPage=True;s.page_setup.fitToWidth=1;s.page_setup.fitToHeight=0
        return s
    metrics=[['全国工作日',"=SUM('全国日历'!E2:E366)"],['法定日数',"=SUM('全国日历'!D2:D366)"],['放假调休窗口天数',"=COUNTIF('全国日历'!C2:C366,\"STATUTORY_HOLIDAY\")+COUNTIF('全国日历'!C2:C366,\"ADJUSTED_REST\")"],['调休上班日',"=COUNTIF('全国日历'!C2:C366,\"MAKEUP_WORKDAY\")"],['公司营业日',"=SUM('营业日历'!B2:B366)"],['旧预排工作日',"=SUM('全国日历'!F2:F366)"],['变更日期数',"=COUNTIF('全国日历'!G2:G366,\"<>0\")"],['项目总数',f"=COUNTA('项目安排'!A2:A{len(plan['projects'])+1})"],['固定已排候选',f'=COUNTIF(\'项目安排\'!F2:F{len(plan["projects"])+1},"PLANNED_FIXED")'],['待客户确认候选',f'=COUNTIF(\'项目安排\'!F2:F{len(plan["projects"])+1},"PROPOSED_CLIENT_CONFIRMATION")'],['阻塞/越范围',f'=COUNTIF(\'项目安排\'!F2:F{len(plan["projects"])+1},"BLOCKED*")+COUNTIF(\'项目安排\'!F2:F{len(plan["projects"])+1},"OUT_OF_SCOPE")'],['分配工序数',f"=COUNTA('工序占用'!A2:A{len(plan['jobs'])+1})"]]
    sheet('看板',['指标','公式结果'],metrics)
    sheet('全国日历',['date','weekday','kind','statutory','national_work','legacy_work','delta','source'],[[r[k] for k in ['date','weekday','kind','statutory','national_work','legacy_work','delta','source']] for r in cal])
    sheet('营业日历',['date','company_open','national_work','local_reason'],[[r[k] for k in ['date','company_open','national_work','local_reason']] for r in company])
    sheet('人员可用',['date','person','role','available','reason'],[[r[k] for k in ['date','person','role','available','reason']] for r in people])
    fields=['id','priority','requested','mode','meeting','status','prep_date','preparer','review_date','reviewer','client_confirmation_required']
    s=sheet('项目安排',fields,[[r[k] for k in fields] for r in plan['projects']]);s.column_dimensions['F'].width=38;s.column_dimensions['K'].width=30
    sheet('工序占用',['project','phase','date','person','role'],[[r[k] for k in ['project','phase','date','person','role']] for r in plan['jobs']])
    sheet('旧新差异',['日期','旧work','新work','差','原因'],[[r['date'],r['legacy_work'],r['national_work'],r['delta'],r['kind']] for r in cal if r['delta']])
    monthly=[]
    for month in range(1,13):
        a=(dt.date(2025,month,1)-dt.date(2025,1,1)).days+2;b=max(i+2 for i,r in enumerate(cal) if r['date'].startswith(f'2025-{month:02}'))
        monthly.append([month,f"=SUM('全国日历'!E{a}:E{b})",f"=SUM('营业日历'!B{a}:B{b})"])
    sheet('月汇总',['月份','全国work','公司open'],monthly)
    sheet('版本与边界',['事项','说明'],[['历史基准','2024-11-18 10:00 Asia/Shanghai；为2025规划，不更改2024考勤/工资。'],['新旧规则','S01 2013版11日；S02签署2024-11-10、公开11-12、2025-01-01施行；年度安排见S03。'],['日历区别','法定13日≠年度放假调休28日≠全国工作日≠公司营业日≠个人可用日。'],['本地约定','S06定义角色容量、优先顺序、严格先准备后复核再会面；不是国家法定SLA。'],['状态','全部为候选内部计划；flexible须客户确认；阻塞披露不等于任务遗漏。'],['限制','不计算加班工资/调休权益，不承诺真实发送/审批/员工排班已生效。'],['复现','运行workflow.py build，验证时使用相同输入与工作目录；哈希不是数字签名/外部批准。']])
    w.save(path)

def build(inp,out,work):
    inp,out,work=inp.resolve(),out.resolve(),work.resolve()
    if inp==out or inp==work or out==work or inp in out.parents or inp in work.parents or out in inp.parents or work in inp.parents or out in work.parents or work in out.parents:raise ValueError('separate nonnested input, output and work paths required')
    out.mkdir(parents=True,exist_ok=True);work.mkdir(parents=True,exist_ok=True)
    if (out/'release_receipt.json').exists():raise ValueError('existing release receipt: build in a fresh output directory')
    initial=snapshots(inp);nodes=[];node_by={}
    def stage(name,data):
        for parent in DAG[name]:
            if parent not in node_by:raise ValueError('dependency not completed: '+parent)
        file=work/(name+'.json');writejson(file,data)
        record=dict(id=name,parents=DAG[name],parent_hashes={p:node_by[p]['node_hash'] for p in DAG[name]},artifact=file.name,artifact_sha256=sha(file),completed_at_utc=dt.datetime.now(dt.timezone.utc).isoformat())
        record['node_hash']=digest(canonical(record));nodes.append(record);node_by[name]=record
    facts,binding=frozen_sources(inp)
    stage('sources',dict(facts=facts,source_freeze=binding))
    legacy_wb=load_workbook(inp/'S05_未审预排v0.xlsx',data_only=True,read_only=True)
    legacy={r[0]:int(r[1]) for r in legacy_wb['v0'].iter_rows(min_row=2,values_only=True)};legacy_wb.close()
    cal=national(facts,legacy);stage('calendar',cal)
    company=service(cal,readcsv(inp/'S09_公司停服.csv'));stage('service',company)
    people=availability(company,readcsv(inp/'S07_岗位表.csv'),readcsv(inp/'S08_请假培训.csv'));stage('availability',people)
    meetings=resolve_meetings(company,readcsv(inp/'S10_客户里程碑.csv'));stage('meetings',meetings)
    plan=allocate(people,meetings);validate_allocation(plan,people);stage('allocation',plan)
    before=work/'spreadsheet';before.mkdir(exist_ok=True);book=before/'operations_plan.xlsx';workbook(book,cal,company,people,plan,facts)
    profile=Path(tempfile.mkdtemp(prefix='go002-office-'))
    p=subprocess.run(['libreoffice','-env:UserInstallation='+profile.as_uri(),'--headless','--convert-to','xlsx','--outdir',str(out),str(book)],capture_output=True,text=True,timeout=120)
    if p.returncode or not (out/'operations_plan.xlsx').exists():raise RuntimeError('spreadsheet recalculation failed')
    scope=['历史规划基准2024-11-18 10:00 Asia/Shanghai。所有客户/员工代码与委托为原创训练情境；以下是待审核通知草稿，未发送。','国家2025节假日采用S02/S03；公司1月26日停服为本地维护，不改变其国家调休上班日性质。2月8日、4月27日、9月28日、10月11日按国家安排上班，仍需核本公司与个人可用性。']
    msgrows=[['项目','请求→候选会面','草稿状态/下一步']]
    labels={'PLANNED_FIXED':'固定会面：内部候选','PROPOSED_CLIENT_CONFIRMATION':'灵活会面：待客户确认','BLOCKED_DATE':'日期阻塞','BLOCKED_CAPACITY':'前置容量不足','OUT_OF_SCOPE':'超出2025范围'}
    for r in plan['projects']:
        status=r['status'];action='内部候选准备'+str(r['prep_date'])+'、复核'+str(r['review_date'])+'；未经复核不对外确认。' if status=='PLANNED_FIXED' else '候选日期须客户确认，未实际变更预约。' if status=='PROPOSED_CLIENT_CONFIRMATION' else '未安排工序；运营负责人补条件/与客户核实，不能假装已完成。'
        msgrows.append([r['id'],r['requested']+' → '+str(r['meeting'] or '未确定'),labels[status]+'；'+action])
    word(out/'client_messages.docx','2025客户里程碑通知草稿',[('口径与发送边界',scope,None),('逐项目草稿',[],msgrows)])
    counts=collections.Counter(r['status'] for r in plan['projects']);stats=dict(national_workdays=sum(r['national_work'] for r in cal),statutory_days=sum(r['statutory'] for r in cal),rest_window_days=sum(r['kind'] in ['STATUTORY_HOLIDAY','ADJUSTED_REST'] for r in cal),company_open_days=sum(r['company_open'] for r in company),status_counts=dict(counts),jobs=len(plan['jobs']))
    note=['S01旧版春节3日、劳动节1日，合计全体公民11日；S02新规则增加除夕与5月2日，自2025-01-01施行为13日。S03年度调休窗口总28日不等于法定假日本身。','旧预排v0是原创未审模板，不是2013办法直接推出的完整年度安排；其缺调休日的错法不可归因于“旧法自动如此”。修订须同步365日历、人员可用、会议提案、两工序占用和所有草稿。','执行依赖：sources→calendar→service→availability/meetings→allocation→documents→preflight。输入变化后旧receipt应失效；不能只改通知维持旧排班或跳过源版本冻结。','已核验的是本地文件一致性与候选约束，不是外部审批/客户同意/人工专业签字。哈希记录不是数字签名，也不保证作者或有权人员作过真实审批。','原文“除个别特殊情形外……一般不超过6天”保留限定，不扩张为任何员工/任何情境绝对禁止；本任务不判断工资、个体假期权益或合法加班。']
    note.insert(0,'2013旧决定落款2013-12-11、自2014-01-01施行；新决定2024-11-08通过、11-10落款、11-12公开，自2025-01-01施行。2025年度安排落款2024-11-12。历史规划基准2024-11-18 10:00 Asia/Shanghai，不把公开日等同起施日，也不追溯更改2024。')
    exceptions=[['项目','状态','处理责任']]+[[r['id'],labels[r['status']],'运营负责人取得2026权威日历再重排。' if r['status']=='OUT_OF_SCOPE' else '运营负责人与客户协商日期/补足前置条件；不得倒借2024工序。' if r['status']=='BLOCKED_CAPACITY' else '客户经理核实固定会面日期；不能自行顺延或确认。'] for r in plan['projects'] if r['status'] not in ['PLANNED_FIXED','PROPOSED_CLIENT_CONFIRMATION']]
    summary=[f"全国工作日{stats['national_workdays']}；公司营业日{stats['company_open_days']}；法定假日{stats['statutory_days']}；年度放假调休窗口{stats['rest_window_days']}。旧预排251工作日，新版248工作日，13个日期的工作标记变化；不是新法增加2日就简单净减2日。",'；'.join(labels[k]+str(v)+'项' for k,v in counts.items())+f"。合计22项目、{stats['jobs']}工序；灵活日期未经客户确认。"]
    word(out/'governance_note.docx','新旧节假日规划迁移及发布复核',[('版本、授权与依赖',note,None),('自动复算摘要',summary,None),('阻塞清单',[],exceptions),('发布顺序',['先由政策/运营负责人核源和本地规则，再由排班负责人核人员及准备—复核—会面依赖；灵活会面由客户经理取得确认，最后由有权负责人批准发布。当前仅为内部复核候选，不触发邮件、预约或真实排班。'],None)])
    own=Path(__file__).resolve()
    if own!=out/'workflow.py':shutil.copy2(own,out/'workflow.py')
    files={n:sha(out/n) for n in OUTPUTS};stage('documents',files)
    if initial!=snapshots(inp):raise ValueError('inputs changed during build')
    validate_allocation(plan,people);stage('preflight',dict(input_stable=True,allocation_constraints_checked=True,external_release=False,stats=stats))
    receipt=dict(schema='go002-release/1',input_hashes=initial,output_hashes=files,stages=nodes,stats=stats,gate='READY_FOR_INTERNAL_REVIEW_WITH_DISCLOSED_BLOCKERS',external_release=False,external_actions_performed=False)
    writejson(out/'release_receipt.json',receipt)
    return receipt

def verify(inp,out,work):
    errors=[]
    try:
        frozen_sources(inp)
        r=json.loads((out/'release_receipt.json').read_text())
        if r['input_hashes']!=snapshots(inp):errors.append('stale_inputs')
        if set(r['output_hashes'])!=set(OUTPUTS):errors.append('output_set')
        for n,h in r['output_hashes'].items():
            if n not in OUTPUTS or sha(out/n)!=h:errors.append('output_drift/'+n)
        if [s['id'] for s in r['stages']]!=list(DAG):errors.append('stage_order_or_missing')
        seen={}
        for item in r['stages']:
            stage=dict(item);recorded=stage.pop('node_hash')
            if digest(canonical(stage))!=recorded:errors.append('node_digest/'+stage['id'])
            if stage['parents']!=DAG[stage['id']] or any(seen.get(p)!=h for p,h in stage['parent_hashes'].items()) or set(stage['parent_hashes'])!=set(stage['parents']):errors.append('dependency/'+stage['id'])
            artifact=stage['artifact']
            if artifact!=stage['id']+'.json' or sha(work/artifact)!=stage['artifact_sha256']:errors.append('artifact/'+stage['id'])
            seen[stage['id']]=recorded
        documents=json.loads((work/'documents.json').read_text());preflight=json.loads((work/'preflight.json').read_text())
        if documents!=r['output_hashes']:errors.append('document_receipt_binding')
        if preflight.get('stats')!=r['stats'] or preflight.get('external_release') is not False:errors.append('preflight_receipt_binding')
        plan=json.loads((work/'allocation.json').read_text());people=json.loads((work/'availability.json').read_text())
        validate_allocation(plan,people)
        if {p['id'] for p in plan['projects']}!={p['id'] for p in readcsv(inp/'S10_客户里程碑.csv')}:errors.append('project_coverage')
        source_payload=json.loads((work/'sources.json').read_text())
        if source_payload['facts']!=json.loads((inp/'S04_官方日历摘录.json').read_text()) or source_payload['source_freeze']!=json.loads((inp/'S11_版本冻结登记.json').read_text()):errors.append('source_payload_binding')
        if r['external_release'] is not False or r['external_actions_performed'] is not False:errors.append('unauthorized_release')
    except (OSError,ValueError,KeyError,TypeError) as e:errors.append('invalid_or_missing_evidence/'+type(e).__name__)
    return dict(valid=not errors,errors=errors,external_release=False)

def main():
    p=argparse.ArgumentParser();p.add_argument('action',choices=['build','verify']);p.add_argument('--input-dir',type=Path,required=True);p.add_argument('--output-dir',type=Path,required=True);p.add_argument('--work-dir',type=Path,required=True);a=p.parse_args()
    if a.action=='build':
        r=build(a.input_dir,a.output_dir,a.work_dir);print(json.dumps(dict(gate=r['gate'],stats=r['stats']),ensure_ascii=False))
    else:
        r=verify(a.input_dir,a.output_dir,a.work_dir);print(json.dumps(r,ensure_ascii=False));raise SystemExit(0 if r['valid'] else 2)

if __name__=='__main__':main()
