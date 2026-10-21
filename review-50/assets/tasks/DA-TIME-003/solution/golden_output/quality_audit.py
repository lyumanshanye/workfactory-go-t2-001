"""Source-grounded TLC batch audit. Internal quality controls, not fare enforcement."""
from pathlib import Path
from collections import Counter
import argparse,csv,datetime,hashlib,json,math,sqlite3,subprocess,tempfile,shutil
import pyarrow.parquet as pq
from openpyxl import Workbook,load_workbook
from openpyxl.styles import Alignment,Font,PatternFill
from openpyxl.utils import get_column_letter
from docx import Document
from docx.shared import Cm,Pt

FILES=['S01_green_2024-12.parquet','S02_green_2025-01.parquet','S03_green_2025-02.parquet','S04_zone_lookup.csv','S05_green_dictionary.pdf','S06_来源与证据.docx','S07_D0.py','S08_R1.json','S09_R2.json','S10_委托权限.docx']
OFFICE=['quality_review.xlsx','intake_memo.docx']
OUTPUTS=['quality_audit.py','quality_ledger.sqlite',*OFFICE,'execution_receipt.json']
FLAGS=['wrong_month','negative_duration','zone_unresolved','zone_not_borough','cbd_missing','distance_nonpositive']
STATES=['structurally_absent','null','negative','zero','positive']
LABELS={'wrong_month':'上客月与源分区不一致','negative_duration':'落客早于上客','zone_unresolved':'zone未匹配','zone_not_borough':'264未知/265市外，不能归城区','cbd_missing':'2025字段存在但值为空','distance_nonpositive':'距离非正','clean':'未触发本合同质量标记'}

def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()

def hashes(inp):
    actual={p.name for p in inp.iterdir() if p.is_file()}
    if actual!=set(FILES):raise ValueError('Input file set incomplete or unexpected')
    return {name:sha(inp/name) for name in FILES}

def contract(inp):
    r=json.loads((inp/'S09_R2.json').read_text())
    if r.get('authority_role')!='数据质量负责人' or r.get('revision') not in ['R2','R3'] or r.get('supersedes') not in ['R1','R2']:raise ValueError('Unapproved contract role/revision')
    if set(r.get('priority',[]))!=set(FLAGS) or len(r['priority'])!=len(FLAGS):raise ValueError('Priority must cover each flag exactly once')
    if type(r.get('quarantine_nonpositive_distance')) is not bool:raise ValueError('Distance disposition switch must be explicit boolean')
    if r.get('timestamp_interpretation')!='America/New_York wallclock; no timezone conversion':raise ValueError('Unsupported timezone contract')
    if {b['month'] for b in r['batches']}!={'2024-12','2025-01','2025-02'} or len(r['batches'])!=3 or {b['file'] for b in r['batches']}!=set(FILES[:3]):raise ValueError('Three distinct batches required')
    return r

def profile(inp,out,r):
    lookup=list(csv.DictReader((inp/FILES[3]).open(encoding='utf-8')))
    zones={int(x['LocationID']):x for x in lookup};assert len(zones)==len(lookup)
    con=sqlite3.connect(out/'quality_ledger.sqlite')
    con.executescript('CREATE TABLE trip_ledger(batch TEXT,row_number INTEGER,pickup_wallclock TEXT,pickup_month TEXT,dropoff_wallclock TEXT,pu INTEGER,do INTEGER,pu_borough TEXT,do_borough TEXT,cbd_state TEXT,cbd_amount REAL,distance REAL,flags_json TEXT,disposition TEXT,PRIMARY KEY(batch,row_number)); CREATE TABLE metadata(key TEXT PRIMARY KEY,value_json TEXT);')
    batches=[];schema=[];cross=Counter();flag_counts=Counter();primary=Counter()
    for b in sorted(r['batches'],key=lambda x:x['month']):
        p=inp/b['file'];f=pq.ParquetFile(p);names=f.schema_arrow.names
        required=['lpep_pickup_datetime','lpep_dropoff_datetime','PULocationID','DOLocationID','trip_distance']
        if not set(required)<=set(names):raise ValueError('Missing operational columns')
        if (b['month']>='2025-01')!=('cbd_congestion_fee' in names):raise ValueError('Unexpected versioned CBD schema')
        for col in names:schema.append(dict(month=b['month'],field=col,type=str(f.schema_arrow.field(col).type)))
        values=Counter();n=0;in_month=0
        for block in f.iter_batches(batch_size=8192):
            insert=[]
            for row in block.to_pylist():
                n+=1;pickup=row['lpep_pickup_datetime'];drop=row['lpep_dropoff_datetime']
                if not isinstance(pickup,datetime.datetime) or not isinstance(drop,datetime.datetime) or pickup.tzinfo is not None or drop.tzinfo is not None:raise ValueError('Missing or timezone-aware timestamp requires explicit review')
                month=pickup.strftime('%Y-%m');cross[(b['month'],month)]+=1;in_month+=int(month==b['month'])
                raw=row.get('cbd_congestion_fee')
                if 'cbd_congestion_fee' not in names:state='structurally_absent'
                elif raw is None:state='null'
                elif not math.isfinite(raw):raise ValueError('Nonfinite CBD requires review, not zero fill')
                else:state='negative' if raw<0 else 'zero' if raw==0 else 'positive'
                values[state]+=1;pu=row['PULocationID'];do=row['DOLocationID'];distance=row['trip_distance']
                flags={'wrong_month':month!=b['month'],'negative_duration':drop<pickup,'zone_unresolved':pu not in zones or do not in zones,'zone_not_borough':pu in [264,265] or do in [264,265],'cbd_missing':state=='null','distance_nonpositive':distance is not None and distance<=0}
                for k,v in flags.items():flag_counts[(b['month'],k)]+=int(v)
                active=[k for k in r['priority'] if flags[k] and (k!='distance_nonpositive' or r['quarantine_nonpositive_distance'])]
                disposition=active[0] if active else 'clean';primary[(b['month'],disposition)]+=1
                insert.append((b['month'],n,pickup.isoformat(),month,drop.isoformat(),pu,do,zones.get(pu,{}).get('Borough'),zones.get(do,{}).get('Borough'),state,raw,distance,json.dumps(flags,sort_keys=True),disposition))
            con.executemany('INSERT INTO trip_ledger VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?)',insert)
        assert n==f.metadata.num_rows
        batches.append(dict(month=b['month'],source_file=b['file'],source_sha256=sha(p),rows=n,in_partition_month=in_month,outside_partition_month=n-in_month,cbd_column_present='cbd_congestion_fee' in names,cbd_states={s:values[s] for s in STATES}))
    summary=dict(batches=batches,total_rows=sum(b['rows'] for b in batches),cross_month=[dict(batch=k[0],pickup_month=k[1],rows=v) for k,v in sorted(cross.items())],flags=[dict(batch=k[0],flag=k[1],rows=v) for k,v in sorted(flag_counts.items())],dispositions=[dict(batch=b['month'],reason=k,rows=primary[(b['month'],k)]) for b in batches for k in FLAGS+['clean']],schema=schema)
    assert con.execute('SELECT COUNT(*) FROM trip_ledger').fetchone()[0]==summary['total_rows']
    for k,v in [('input_sha256',hashes(inp)),('contract',r),('summary',summary)]:con.execute('INSERT INTO metadata VALUES(?,?)',(k,json.dumps(v,ensure_ascii=False,sort_keys=True)))
    con.commit();assert con.execute('PRAGMA integrity_check').fetchone()[0]=='ok';con.close();return summary

def workbook(out,s,r):
    w=Workbook();w.remove(w.active)
    def sheet(name,rows,widths):
        q=w.create_sheet(name)
        for row in rows:q.append(row)
        for c in q[1]:c.font=Font(name='Noto Sans CJK SC',bold=True,color='FFFFFF',size=10);c.fill=PatternFill('solid',fgColor='173C45')
        for i,v in enumerate(widths,1):q.column_dimensions[get_column_letter(i)].width=v
        for row in q:
            q.row_dimensions[row[0].row].height=34 if row[0].row==1 else 29
            for c in row:c.alignment=Alignment(wrap_text=True,vertical='center',horizontal='left' if isinstance(c.value,str) and not c.value.startswith('=') else 'center');c.font=Font(name='Noto Sans CJK SC',size=10,bold=c.row==1,color='FFFFFF' if c.row==1 else '173C45')
        q.freeze_panes='A2';q.auto_filter.ref=q.dimensions;q.print_title_rows='1:1';q.sheet_properties.pageSetUpPr.fitToPage=True;q.page_setup.orientation='landscape';q.page_setup.paperSize=q.PAPERSIZE_A3;q.page_setup.fitToWidth=1;q.page_setup.fitToHeight=1
        return q
    rows=[['源分区月','原始行','上客本月','上客跨月','SQLite行','唯一处置计数','行守恒差','结论']]
    for i,b in enumerate(s['batches'],2):rows.append([b['month'],b['rows'],b['in_partition_month'],b['outside_partition_month'],b['rows'],sum(v['rows'] for v in s['dispositions'] if v['batch']==b['month']),f'=B{i}-C{i}-D{i}',r['revision']+'：全量保留，待核问题见后页'])
    rows.append(['合计','=SUM(B2:B4)','=SUM(C2:C4)','=SUM(D2:D4)','=SUM(E2:E4)','=SUM(F2:F4)','=SUM(G2:G4)','只证明所交三源处理覆盖，非官方完整性保证'])
    q=sheet('批次行守恒',rows,[17,16,16,16,16,19,17,60]);q.row_dimensions[5].height=55
    sheet('上客月归属',[['源分区月','真实上客月','行数','解释']]+[[x['batch'],x['pickup_month'],x['rows'],'原始墙钟未转UTC；不同批次相同月份不去重或自动移入另一源文件'] for x in s['cross_month']],[20,20,18,100])
    fields=sorted({x['field'] for x in s['schema']});schema={(x['month'],x['field']):x['type'] for x in s['schema']}
    def boundary(c):
        if c=='cbd_congestion_fee':return '结构不存在≠null≠数值0；不同于congestion_surcharge'
        if c in ['lpep_pickup_datetime','lpep_dropoff_datetime']:return '无时区时间戳按给定纽约墙钟；不作UTC转换'
        if c in ['PULocationID','DOLocationID']:return '原zone ID；264未知/265市外不当普通城区'
        if c=='congestion_surcharge':return 'NYS费项，不代填cbd_congestion_fee'
        if c in ['tip_amount','total_amount']:return '按官方字典不含现金小费，不推断全部收入'
        return '保留官方原字段与类型；本审阅不作逐程计费合法性判断'
    sheet('Schema版本',[['字段','2024-12','2025-01','2025-02','使用边界']]+[[c,*[schema.get((m,c),'结构不存在') for m in ['2024-12','2025-01','2025-02']],boundary(c)] for c in fields],[34,23,23,23,74])
    sheet('CBD五态',[['源分区月','无此字段','有列但null','负值','零值','正值','核算合计','与总行差']]+[[b['month'],*[b['cbd_states'][k] for k in STATES],f'=SUM(B{i}:F{i})',f'=G{i}-\'批次行守恒\'!B{i}'] for i,b in enumerate(s['batches'],2)],[20,19,19,19,19,19,20,20])
    flags={(x['batch'],x['flag']):x['rows'] for x in s['flags']};primary={(x['batch'],x['reason']):x['rows'] for x in s['dispositions']}
    sheet('异常与唯一处置',[['月','标记/处置原因','多标签行数','唯一处置行数','区别']]+[[b['month'],LABELS[k],flags.get((b['month'],k),None),primary.get((b['month'],k),0),'多标签可重叠，不能相加当异常总量；唯一处置按合同优先级，含clean守恒'] for b in s['batches'] for k in FLAGS+['clean']],[18,39,22,24,85])
    recommendations=[['问题/责任','已知事实和动作'],['接入决定','仅内部质量待审；原始层可按SHA归档，业务月报层暂不自动放行。数据质量负责人审阅隔离规则后批准，不能宣称TLC官方已验真。'],['全量覆盖','三分区共'+str(s['total_rows'])+'行全部入SQLite。行号从各源物理读取顺序1开始，(batch,row_number)唯一；不以抽样冒全量。'],['日期/版本','字典2025-03-18说明CBD对应2025-01-05起费用；不是原文件首次发布时间。业务回放2026-10-21 09:00 Asia/Shanghai，真实执行UTC另记。'],['未解决事项','源提供商可有误，TLC不担保准确性；264为未知、265为市外，join成功不等可归城区。落客早于上客等只提示补证，不推断违法/造假。'],['缺值和费项','2024无CBD列不造零；2025 null不认免费；congestion_surcharge与cbd_congestion_fee不同，不互相补值或并为同一政策。现金小费不在tip_amount/total中，不由字段推断总收入。'],['下一步','工程师交隔离账与复现命令，业务负责人确认分析分母，数据质量负责人审核；不得删除异常、补造金额、自动上线或宣称政策因果效应。'],['复现与边界','python quality_audit.py build --input-dir /app/input_files --output-dir 新空目录（先复制本脚本）；verify同参数。核输入/脚本/3成果SHA，合法变更新目录重建，不手改旧回执。']]
    q=sheet('接入建议',recommendations,[27,150])
    for row in range(2,9):q.row_dimensions[row].height=68
    w.properties.creator='Work50';w.save(out/OFFICE[0]);w.close()
    with tempfile.TemporaryDirectory(prefix='tlc-recalc-') as tmp:
        t=Path(tmp);dest=t/'recalculated';dest.mkdir()
        run=subprocess.run(['libreoffice','-env:UserInstallation='+(t/'profile').as_uri(),'--headless','--convert-to','xlsx','--outdir',str(dest),str((out/OFFICE[0]).resolve())],capture_output=True,text=True,timeout=120)
        if run.returncode or not (dest/OFFICE[0]).exists():raise RuntimeError('Actual formula recalculation failed: '+run.stdout+run.stderr)
        shutil.copy2(dest/OFFICE[0],out/OFFICE[0])
    w=load_workbook(out/OFFICE[0],data_only=True)
    assert w['批次行守恒']['B5'].value==s['total_rows']
    assert all(w['批次行守恒'][f'G{i}'].value==0 for i in range(2,6))
    assert all(w['CBD五态'][f'H{i}'].value==0 for i in range(2,5));w.close()

def memo(out,s,r):
    d=Document();sec=d.sections[0];sec.top_margin=sec.bottom_margin=Cm(1.6);d.styles['Normal'].font.name='Noto Sans CJK SC';d.styles['Normal'].font.size=Pt(10)
    d.core_properties.author='Work50';d.add_heading('跨年Green数据接入｜质量待审',0)
    d.add_paragraph('致数据质量负责人和月报业务负责人｜原创内部业务回放2026-10-21 09:00 Asia/Shanghai；不是TLC官方审批或历史抓取证明。执行时刻及冻结SHA见execution_receipt.json。')
    d.add_heading('一、结论与可用范围',1);d.add_paragraph('暂不自动放行业务月报；三份真实源共'+str(s['total_rows'])+'行全量保留到quality_ledger.sqlite。原始层可按哈希归档；业务层须先核异常/分母与版本口径，不删除异常、不补造金额或推断政策因果效应。TLC声明数据来自技术提供商，不担保准确性，作者也不能以校验通过替官方验真。')
    for b in s['batches']:
        f={x['flag']:x['rows'] for x in s['flags'] if x['batch']==b['month']}
        d.add_paragraph(b['month']+'：'+str(b['rows'])+'行；上客本月'+str(b['in_partition_month'])+'、跨月'+str(b['outside_partition_month'])+'；落客早于上客'+str(f['negative_duration'])+'。CBD五态（无列/null/负/零/正）='+ '/'.join(str(b['cbd_states'][k]) for k in STATES)+'。这些是源分区口径，不是重分月后的去重行程量。')
    d.add_heading('二、版本、消息与时间',1);d.add_paragraph('采用负责人的'+r['revision']+'，保留R1历史，不采纳较晚营销删异常/填零建议。有权R2增加全量行账、schema五态及唯一处置；这属于原创内部合同，不是TLC官方质量标准。无时区时间戳按America/New_York当地墙钟，不作UTC转换。真实抓取2026，不能伪称此快照2025已存在。')
    d.add_paragraph('官方网页称2025起新增cbd_congestion_fee；字典2025-03-18说明费用从2025-01-05开始。旧文件无该列不等金额零，新文件null也不等免费；congestion_surcharge是另一费项，不能替代。只有字段与日期不足以逐程认定计费违法；不据前后月份差异声称政策有效。')
    d.add_page_break();d.add_heading('三、行守恒与可追溯账',1)
    d.add_paragraph('quality_ledger.sqlite以(batch,row_number)唯一保留每源每行，上客/落客墙钟、zone、CBD状态/原值、距离、全部质量标记和唯一处置。每批原行数=SQLite行数=唯一处置之和；本月+跨月=原行数。多标签同一行可重叠，标签合计不是异常总行数。优先级：'+' → '.join(r['priority'])+' → clean；距离是否隔离为'+str(r['quarantine_nonpositive_distance'])+'，不改变原始距离/标记。')
    d.add_paragraph('Schema版本页逐字段列实际类型/缺列；CBD五态页保留结构缺失、null、负、零、正且总数守恒。zone 264是Unknown，265是Outside of NYC，不是普通城区；即使能join也不能直接计入区级空间覆盖。原日期异常和非正距离只标记待核，不认定驾驶人、提供商违法或造假。')
    d.add_heading('四、D0错误与后续动作',1);d.add_paragraph('D0为原创有缺陷旧稿：只读1000行冒全量、文件名代上客月、缺列/null全填0、过滤非正距离。新版实际遍历三份全部行，分开两种月份，五态不补值，保留异常。无修改任何官方输入。')
    d.add_paragraph('工程师向数据质量负责人交带行号的隔离账和当前哈希；业务负责人确认区级报表分母、跨月归属和特殊zone处置；必要时联系来源提供方核实日期/缺值。待负责人核准再建下游业务视图，不从原始层删行。未做真实联系、官方确认、自动发布、模型能力测试或专家签字。')
    d.add_heading('五、复现与失败恢复',1);d.add_paragraph('依赖Python、pyarrow、openpyxl、python-docx、sqlite3及LibreOffice（真实重算公式缓存）。新建空输出目录并复制quality_audit.py，执行python quality_audit.py build --input-dir /app/input_files --output-dir 新目录；随后verify同参数。build拒不完整输入、无权合同与未知schema；verify核输入/当前脚本/SQLite/两Office哈希及行守恒。变更后旧回执失效，在新目录重建；哈希只是字节一致证明，不是可信签章或全部语义正确证明。')
    d.add_paragraph('来源：S01–S03官方链接的完整Parquet；S04官方zone表；S05官方2页字典；S06含原站免责声明/取证日期。S07–S10的旧稿、合同、权限和业务日期均原创演练，不代表市政数据生产指令。')
    d.save(out/OFFICE[1])

def build(inp,out):
    before=hashes(inp);r=contract(inp);out.mkdir(parents=True,exist_ok=True)
    if {p.name for p in out.iterdir()}!={'quality_audit.py'}:raise ValueError('Use fresh output directory containing only this script')
    if sha(out/'quality_audit.py')!=sha(Path(__file__)):raise ValueError('Output script not this executed version')
    s=profile(inp,out,r);workbook(out,s,r);memo(out,s,r)
    assert before==hashes(inp)
    result=dict(task_id='DA-TIME-003',execution_at_utc=datetime.datetime.now(datetime.timezone.utc).isoformat(),input_sha256=before,script_sha256=sha(Path(__file__)),contract_revision=r['revision'],contract_authority=r['authority_role'],summary=s,output_sha256={n:sha(out/n) for n in ['quality_ledger.sqlite',*OFFICE]},release_status='INTERNAL_REVIEW_REQUIRED',external_publish_performed=False,scope='Actual full-row audit of supplied frozen batches, not source accuracy guarantee, fare enforcement, causal inference, expert approval or trusted attestation')
    (out/'execution_receipt.json').write_text(json.dumps(result,ensure_ascii=False,indent=2));return result

def verify(inp,out):
    e=json.loads((out/'execution_receipt.json').read_text());assert set(p.name for p in out.iterdir())==set(OUTPUTS)
    assert e['input_sha256']==hashes(inp);assert e['script_sha256']==sha(Path(__file__))==sha(out/'quality_audit.py')
    assert e['contract_revision']==contract(inp)['revision'] and not e['external_publish_performed'] and e['release_status']=='INTERNAL_REVIEW_REQUIRED'
    assert e['output_sha256']=={n:sha(out/n) for n in ['quality_ledger.sqlite',*OFFICE]}
    c=sqlite3.connect('file:'+str((out/'quality_ledger.sqlite').resolve())+'?mode=ro',uri=True)
    assert c.execute('PRAGMA integrity_check').fetchone()[0]=='ok'
    assert json.loads(c.execute('SELECT value_json FROM metadata WHERE key="summary"').fetchone()[0])==e['summary']
    assert json.loads(c.execute('SELECT value_json FROM metadata WHERE key="input_sha256"').fetchone()[0])==hashes(inp)
    for b in e['summary']['batches']:
        assert c.execute('SELECT COUNT(*) FROM trip_ledger WHERE batch=?',(b['month'],)).fetchone()[0]==b['rows']==pq.ParquetFile(inp/b['source_file']).metadata.num_rows
        assert b['in_partition_month']+b['outside_partition_month']==b['rows']==sum(b['cbd_states'].values())
    c.close();return dict(task_id='DA-TIME-003',verified=True,total_rows=e['summary']['total_rows'],release_status=e['release_status'])

def main():
    p=argparse.ArgumentParser();p.add_argument('command',choices=['build','verify']);p.add_argument('--input-dir',type=Path,required=True);p.add_argument('--output-dir',type=Path,required=True);a=p.parse_args()
    e=build(a.input_dir,a.output_dir) if a.command=='build' else verify(a.input_dir,a.output_dir)
    print(json.dumps({k:v for k,v in e.items() if k not in ['summary','input_sha256','output_sha256']},ensure_ascii=False))

if __name__=='__main__':main()
