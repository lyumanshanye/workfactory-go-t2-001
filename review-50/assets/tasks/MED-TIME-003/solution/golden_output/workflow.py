"""Frozen ethics-training workflow; no real ethics approval or system upload."""
from pathlib import Path
import argparse,datetime,hashlib,json,subprocess,tempfile
from docx import Document
from docx.shared import Cm,Pt as WPt
from openpyxl import Workbook,load_workbook
from openpyxl.styles import Font,PatternFill,Alignment
from pptx import Presentation
from pptx.util import Inches,Pt
from pptx.dml.color import RGBColor

NAMES=['S01_2016公报节选.pdf','S02_2023公报节选.pdf','S03_官方解读第七问.docx','S04_教学案例R1.json','S05_权限与消息.docx','S06_R2补充.json','S07_D0培训稿.docx','S08_工作流合同.json','S09_培训委托.docx','S10_来源登记.json']
OUTPUTS=['training_matrix.xlsx','trainer_brief.pptx','facilitator_notes.docx']
DEPS={'S0':[],'S1':['S0'],'S2':['S0'],'S3':['S1','S2'],'S4':['S1'],'S5':['S3','S4'],'S6':['S3','S4'],'S7':['S3','S4'],'S8':['S5','S6','S7'],'S9':['S8']}
RULES=[
 ['T1','两办法并行','2016十一号令为部门规章；2023统一制度指导实践，医疗机构违规可据11号令处理，不写全面废止。','S03第七问；S01第45–49条'],
 ['T2','适用主体','2023覆盖医疗卫生机构、高等学校、科研院所等；不是仅医院，也不因学生作业自动排除。','S02第2–3条'],
 ['T3','一般/简易','一般会议审查；符合简易条件可由主任委员指定两个或以上委员，意见不一致等须转会议。','S02第15/31条；S01第24条仅作旧规范对照'],
 ['T4','免审不是免同意','第32条有无伤害/无敏感个人信息或商业利益等前提；匿名化不等去标识化。2016医疗机构适用范围内的免除签署制度与免除伦理审查不同，不能由32条自行免签。','S02第32/33条；S01第39条（医疗机构范围）'],
 ['T5','多机构','牵头与参与机构均应组织伦理审查，参与机构还需跟踪；可依法采用简易确认，不直接拿牵头意见替代。','S02第27/31(四)条'],
 ['T6','委托','书面委托有能力的机构或区域委员会；医疗机构还需不低于本机构等级的医疗委员会或区域委员会，受托须跟踪。','S02第5/14条'],
 ['T7','知情同意','限制/无民事行为能力者需监护人书面同意，同时在参与者可理解范围告知并征得其同意，不能只收签字。','S02第34条'],
 ['T8','时间和上传','受理后30天内开展审查并出意见；紧急情况及时，疫情突发一般72小时，不降低质量；获批研究实施前各主体上传并及时更新。','S02第16/24条'],
 ['T9','备案起点','新设委员会设立后3个月备案；既设委员会过渡6个月、既获批研究9个月分别按第54条起点，不混为一条。年度报告每年3月31日前。','S02第13/54条；S01第14/50条'],
 ['T10','变化再审','批准后修改方案/同意书/招募等须提交审查；未授权同事无权替委员会批准或覆盖负责人受控补充。','S02第23/30条；S05/S06原创权限'],
]
LABELS={'RISK_INFO_HOLD':'补风险资料，未定程序','EXEMPTION_CANDIDATE':'免审候选，交机构核定','ANONYMITY_NOT_ESTABLISHED':'未证匿名化，不直接免审','SIMPLE_CANDIDATE':'简易候选，至少两委员','MEETING_REVIEW':'转会议审查','LOCAL_REVIEW_REQUIRED':'本机构审查并跟踪','WRITTEN_DELEGATION':'书面委托有能力委员会','DELEGATION_SCOPE_HOLD':'委托等级不符，重新选择','CONSENT_SUPPLEMENT':'补参与者告知与同意','REVIEW_NOT_EXEMPT':'商业利益，不按32条免审','LEGACY_COMMITTEE':'既设委员会过渡备案','NEW_COMMITTEE':'新设委员会三个月备案'}

def digest(p):return hashlib.sha256(Path(p).read_bytes()).hexdigest()
def dump(p,data):Path(p).write_text(json.dumps(data,ensure_ascii=False,indent=2)+'\n')
def load(p):return json.loads(Path(p).read_text())
def inputs(root):return {name:digest(root/name) for name in NAMES}


def decide(c):
    kind=c['kind'];owner='机构伦理审查办公室会同研究负责人'
    if kind=='new_application':
        route='RISK_INFO_HOLD';action='补风险与科学性等申请材料，由委员会判断程序，秘书不得从无干预推定最小风险。';basis='S02第18/19/31条'
    elif kind=='data':
        if c['commercial_interest']:
            route='REVIEW_NOT_EXEMPT';action='不能仅因公开/匿名主张第32条免审；提交机构按事实核定审查路径。'
        elif c.get('reidentification_key_retained'):
            route='ANONYMITY_NOT_ESTABLISHED';action='保留映射钥匙不能当已匿名；补数据可识别性/来源/用途资料，不直接免审。'
        elif c['lawfully_public'] and c['no_harm'] and not c['sensitive_personal_information']:
            route='EXEMPTION_CANDIDATE';action='给定前提可提出32(一)免审候选，按本团队流程交机构核定；不是秘书已作伦理批准。'
        else:
            route='RISK_INFO_HOLD';action='补第32条所需前提证据，不凭标签认定免审。'
        basis='S02第32条；S08原创内部核定流程'
        if c.get('reidentification_key_retained'):basis+='；S03术语旁注个保法第73条'
    elif kind=='simple_review':
        if c.get('reviewer_disagreement'):
            route='MEETING_REVIEW';action='简易中委员意见不一致，转会议审查；不得按多数或主任单签强过。'
        else:
            assert c['approved_existing'] and c['minor_amendment'] and c['risk_benefit_unchanged']
            route='SIMPLE_CANDIDATE';action='符合给定小改/不影响风险受益条件，可申请简易，由主任指定两个或以上委员并向会议报告。'
        basis='S02第23/31条'
    elif kind=='multicenter':
        route='LOCAL_REVIEW_REQUIRED';action='牵头已批不代本机构审查/跟踪；可依法走参与机构简易确认，不能直接开研。';basis='S02第27/31(四)条'
    elif kind=='delegation':
        if c['institution_type']=='hospital' and c['proposed_committee_type']=='hospital' and c['proposed_grade']<c['institution_grade']:
            route='DELEGATION_SCOPE_HOLD';action='拟受托医院等级低于委托方，改选不低等级的有能力医疗委员会或区域委员会；书面委托、跟踪及本机构设委责任均待落实。'
        else:
            route='WRITTEN_DELEGATION';action='独立企业可书面委托有能力机构委员会或区域委员会，受托需跟踪；不编企业医院等级要求或伪造已受托。'
        basis='S02第5/14条；S03第七问的机构/监管边界'
    elif kind=='consent':
        route='CONSENT_SUPPLEMENT';action='除监护人书面同意外，还须在参与者可理解范围告知并征得其同意；本题未完成，不宣布可入组。';basis='S02第34条'
    elif kind=='committee':
        if c['established_on']<'2023-02-18':
            route='LEGACY_COMMITTEE';action='既设委员会按54条过渡6个月办理并上传，不套新设3个月；不在未核发布起点时伪算精确截止日。'
        elif c['established_on']>='2023-02-28':
            route='NEW_COMMITTEE';action='新设委员会按设立起3个月备案/上传；医疗向执业登记机关，其他按隶属向上级主管；年度报告另看3月31日。'
        else:
            raise ValueError('Establishment falls between signature and known public explanation; confirm authoritative publication date before assigning transition')
        basis='S02第13/54条'
    else:raise ValueError('Unsupported kind: '+str(kind))
    fact=c['teaching_fact']+(' 本轮采用commercial_interest='+str(c['commercial_interest']).lower()+'。' if kind=='data' else '')
    return dict(case_id=c['case_id'],title=c['title'],fact=fact,route=route,label=LABELS[route],basis=basis,action=action,owner=owner,ethics_approved=False,study_start_authorized=False)


def book(path,rows):
    w=Workbook();w.remove(w.active)
    tabs={'规则对照':[['主题','要点','培训口径','依据']]+RULES,
          '案例讲解':[['案例','题名','状态','处理','依据','下一责任接口','允许开研']]+[[r['case_id'],r['title'],r['route'],r['action'],r['basis'],r['owner'],'否'] for r in rows],
          'D0更正':[['旧稿错误','正确处理','定位'],['2023已废止2016','并行且处罚依据区分','T1/S03七'],['主任一人即可简易','至少两个委员，不一致转会议','T3/S02第31条'],['去标识化/商业公开数据都免审','前提逐一核，不把钥匙保留当匿名','T4/C/J'],['牵头已批全点直接用','本机构审查/跟踪，可依法简易确认','T5/F'],['监护人签字就可入组','参与者理解范围内还须告知并征同意','T7/I'],['所有委员会3个月','分新设/过渡6个月，已批研究9个月另列','T9/K/L']],
          '队列统计':[['项目','数量','说明'],['案例数',"=COUNTA('案例讲解'!A2:A13)",'12个原创教学情境'],['免审候选','=COUNTIF(\'案例讲解\'!C2:C13,"EXEMPTION_CANDIDATE")','仅候选，不是实际批准'],['简易候选','=COUNTIF(\'案例讲解\'!C2:C13,"SIMPLE_CANDIDATE")','至少两委员'],['允许开研','=COUNTIF(\'案例讲解\'!G2:G13,"是")','应为0'] ]}
    widths={'规则对照':[12,24,95,52],'案例讲解':[9,28,37,90,42,34,14],'D0更正':[42,72,30],'队列统计':[26,18,50]}
    for name,data in tabs.items():
        s=w.create_sheet(name)
        for row in data:s.append(row)
        for i,c in enumerate(s[1],1):c.font=Font(bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='173C45')
        for i,size in enumerate(widths[name],1):s.column_dimensions[chr(64+i)].width=size
        for row in s:
            s.row_dimensions[row[0].row].height=78 if name=='案例讲解' else 62
            for c in row:c.alignment=Alignment(wrap_text=True,vertical='center',horizontal='center' if name=='队列统计' and c.column==2 else None)
        s.freeze_panes='A2';s.print_title_rows='1:1';s.print_options.horizontalCentered=True;s.sheet_properties.pageSetUpPr.fitToPage=True;s.page_setup.orientation='landscape';s.page_setup.paperSize=s.PAPERSIZE_A3;s.page_setup.scale=None;s.page_setup.fitToWidth=1;s.page_setup.fitToHeight=1
    # Real Calc evaluation, not manually inserted cached numbers.
    with tempfile.TemporaryDirectory(prefix='med003-calc-') as temp:
        folder=Path(temp);source=folder/'training_matrix.xlsx';out=folder/'out';out.mkdir();w.save(source)
        p=subprocess.run(['libreoffice','-env:UserInstallation='+(folder/'profile').as_uri(),'--headless','--convert-to','xlsx','--outdir',str(out),str(source)],capture_output=True,text=True,timeout=90)
        assert p.returncode==0 and (out/source.name).is_file(),p.stdout+p.stderr
        path.write_bytes((out/source.name).read_bytes())


def slides(path,rows):
    p=Presentation();p.slide_width=Inches(13.333);p.slide_height=Inches(7.5)
    def slide(title,blocks,footer):
        s=p.slides.add_slide(p.slide_layouts[6]);s.background.fill.solid();s.background.fill.fore_color.rgb=RGBColor.from_string('F1F5F6')
        def text(x,y,w,h,value,size=20,color='173C45'):
            box=s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h));tf=box.text_frame;tf.word_wrap=True
            for i,line in enumerate(value.split('\n')):
                q=tf.paragraphs[0] if i==0 else tf.add_paragraph();q.text=line;q.font.size=Pt(size);q.font.name='Noto Sans CJK SC';q.font.color.rgb=RGBColor.from_string(color)
        text(.6,.35,12.1,.8,title,29)
        for i,block in enumerate(blocks):text(.65,1.55+i*2.25,12,2.05,block,19)
        text(.65,6.95,12,.35,footer,11,'48636B')
    slide('伦理培训更新｜两规范并行，不能见新废旧',['受众：医疗研究人员、学生、科研管理/伦理办公室\n2023-03-01 10:00 Asia/Shanghai｜原创教学情境；待机构审阅','本课是培训材料和内部核查路线，不是伦理批准、入组许可或系统已上传。'], 'S01/S02官方公报；S03官方索引转录；源抓取2026-10-20不是业务时点')
    slide('版本、机构与时间：先问依据和权限',['2016十一号令：2016-10-12公布，2016-12-01起施。\n2023四部门办法：落款2月18日，自发布起施；2月27日是本题官方解读发布日期，不冒称通知准确发布时刻。','官方解读第七问明确并行：实践可按2023办法指导；医疗违规可据11号令处理。\n3月1日负责人R2补充有效，较晚营销便笺无权覆盖。'], 'S01印刷1/7页；S02第54条；S03七；S05/S06原创权限')
    slide('三条常混边界，先讲清再做案例',['简易≠免审：主任指定两个或以上委员，分歧等转会议。\n免审≠免除签署同意：32条先核无伤害/无敏感个人信息或商业利益等前提。','牵头意见≠参与点免程序；秘书路线建议≠委员会决定。\n新设3个月、既设过渡6个月、既批研究9个月不是一个起点。'], 'S02第13/27/31/32/33/54条；S01第39条（医疗机构范围）；S08仅内部核定')
    for i in range(0,len(rows),2):
        blocks=[]
        for r in rows[i:i+2]:blocks.append(r['case_id']+'｜'+r['title']+'\n'+r['fact']+'\n'+r['route']+'｜'+r['action'])
        slide('案例讨论 '+rows[i]['case_id']+' / '+rows[i+1]['case_id'],blocks,'；'.join(r['case_id']+': '+r['basis'] for r in rows[i:i+2]))
    j=next(r for r in rows if r['case_id']=='J')
    slide('讲师收束：更新必须穿透全部产物',['先冻结源 → 并行核规则/有权补充 → 合并案例 → 课件/底稿/讲师说明 → 勾稽 → 待审回执。\nJ按受控'+j['adopted_revision']+'：'+j['label']+'，不能静默沿用旧事实。','出现新补充：旧快照/节点/产物即过期，重跑并核全部SHA。\n机构伦理办公室与负责人审阅，未完成真实审查/备案/上传；不上线、不发给参与者。'], 'S08原创依赖合同；S05/S06；workflow.py与build_receipt.json实际执行证据')
    p.save(path)


def notes(path,rows):
    d=Document();s=d.sections[0];s.top_margin=s.bottom_margin=Cm(1.7);d.styles['Normal'].font.name='Noto Sans CJK SC';d.styles['Normal'].font.size=WPt(10)
    d.add_heading('讲师与办公室交接｜待机构审阅',0)
    d.add_paragraph('基准2023-03-01 10:00 Asia/Shanghai。面向研究人员、学生及科研管理人员；业务/案例/SOP原创，政策来自真实冻结来源。12情境均未取得伦理批准或开研授权。')
    d.add_heading('一、讲前必须纠正的旧稿',1)
    for r in RULES:d.add_paragraph(r[0]+' '+r[1]+'：'+r[2]+'（'+r[3]+'）')
    d.add_page_break();d.add_heading('二、讨论组织与关键追问',1)
    d.add_paragraph('先由学员判断“缺证/候选/需组织审查”，再追问依据、缺口和下一责任人。委员会依法决定，秘书只汇总/核形式。免审候选交机构核定是S08内部培训操作，不假称条文直接授予秘书此权。')
    for r in rows:d.add_paragraph(r['case_id']+'｜'+r['label']+'：'+r['action'])
    d.add_page_break();d.add_heading('三、真实依赖工作流和范围',1)
    d.add_paragraph('重建命令：python3 workflow.py build --input-dir /app/input_files --output-dir /app/output --work-dir /tmp/med003-new-run。使用一个尚无state.json的新工作目录，保留旧运行证据。验证：python3 workflow.py verify --input-dir /app/input_files --output-dir /app/output --work-dir /tmp/med003-new-run。stage S5等只在前置完成且哈希一致时可运行。')
    j=next(r for r in rows if r['case_id']=='J')
    d.add_paragraph('S0冻结10输入；S1规则与S2受控事实并行依赖S0；S3须两者合并；S4教学口径依赖规则；S5/S6/S7分别生成工作簿/课件/讲师说明；S8勾稽三产物；S9只签发内部待审回执。J按有权'+j['adopted_revision']+'为“'+j['label']+'”，晚到无权便笺不能覆盖。快照/节点/产物任一被改，verify必须拒绝；重新构建不得沿用旧receipt。')
    d.add_paragraph('变更责任：培训负责人核业务补充，伦理办公室核法规/程序，讲师同步全部课件，研究负责人补实际审查材料。后台备案/上传、临床安全、机构执业/能力、参与者知情同意实际取得与专家签字均未执行。未给完整事实不裁定临床风险/适当性，不伪造审批。')
    d.add_paragraph('时点：2016公布10月12日/起施12月1日；2023通知落款2月18日，自发布起施；2月27日为官方解读发布日期，不据此冒称通知准确发布时刻。不把公报下载地址2023/202305或后期上传文件名当法律起算；具体过渡截止日未要求、未伪算。K在2月1日设立、L在3月1日设立，分别位于前/后明确范围；若新案例落在落款与已知公开解读之间，应转核准确起点而不是机械分类。')
    d.add_paragraph('S01/S02为官方公报页面节选，保留原印刷页码；S02首提取页上方可能含前一条公报尾段，不属于本办法。S03为官方索引原文转录而非已下载原HTML，直连412已披露。取证2026-10-20，不是2023现场记录或授权。')
    d.save(path)


def check_office(output,rows):
    w=load_workbook(output/OUTPUTS[0],data_only=True);s=w['案例讲解']
    actual={s.cell(i,1).value:[s.cell(i,j).value for j in (3,4,5,6,7)] for i in range(2,s.max_row+1)}
    expected={r['case_id']:[r['route'],r['action'],r['basis'],r['owner'],'否'] for r in rows}
    assert actual==expected,'Workbook routes/actions/owners stale or modified'
    assert [w['队列统计'].cell(i,2).value for i in range(2,6)]==[12,sum(r['route']=='EXEMPTION_CANDIDATE' for r in rows),sum(r['route']=='SIMPLE_CANDIDATE' for r in rows),0]
    w.close();p=Presentation(output/OUTPUTS[1]);blocks=[s.text for slide in p.slides for s in slide.shapes if s.has_text_frame]
    assert len(p.slides)==10
    for r in rows:
        matches=[b for b in blocks if r['case_id']+'｜'+r['title'] in b]
        assert len(matches)==1 and r['route'] in matches[0] and r['action'] in matches[0],'Stale case-specific slide'
    d=Document(output/OUTPUTS[2]);body='\n'.join(p.text for p in d.paragraphs)
    for r in rows:assert r['case_id']+'｜'+r['label'] in body and r['action'] in body
    assert '并行' in body and '未取得伦理批准' in body
    return dict(cases=12,slide_count=len(p.slides),formula_cached_values=[12,sum(r['route']=='EXEMPTION_CANDIDATE' for r in rows),1,0],study_start_authorized=False)


def stage(node,root,output,work):
    current=inputs(root);statepath=work/'state.json'
    if statepath.exists():
        state=load(statepath);assert state['input_sha256']==current,'Input changed; initialize a new run, never reuse stale nodes'
        assert state['workflow_sha256']==digest(__file__),'Workflow changed; new run required'
        for prior,entry in state['stages'].items():
            assert digest(work/entry['artifact'])==entry['sha256'],'Completed intermediate changed: '+prior
    else:
        assert node=='S0','Missing dependency: S0 not completed'
        state=dict(input_sha256=current,workflow_sha256=digest(__file__),stages={})
    contract=load(root/NAMES[7]);assert contract['dependencies']==DEPS
    assert node in DEPS
    for dep in DEPS[node]:
        assert dep in state['stages'],'Missing dependency '+dep
        entry=state['stages'][dep]
        assert digest(work/entry['artifact'])==entry['sha256'],'Dependency artifact changed: '+dep
    assert node not in state['stages'],'Already completed; use a fresh run instead of overwriting provenance'
    if node=='S0':data=dict(inputs=current,nature='actual source byte snapshot, not institutional verification')
    elif node=='S1':
        registry=load(root/NAMES[9])
        for selection in registry['source_selections']:
            assert digest(root/selection['input_file'])==selection['input_sha256'],'Policy selection changed; source review and registry update required'
        assert digest(root/NAMES[2])==registry['faq']['transcription_input_sha256'],'Official explanation transcription changed'
        data=dict(rules=RULES,parallel_regulations=True,source_note='Manually interpreted frozen provisions, not automatic legal text understanding')
    elif node=='S2':
        packet=load(root/NAMES[3]);change=load(root/NAMES[5]);assert change['authorized_role']==contract['change_authority'] and change['target_packet']=='R1'
        assert change['issued_at']<packet['as_of']
        by={c['case_id']:c for c in packet['cases']};assert len(by)==12
        for edit in change['changes']:
            assert by[edit['case_id']][edit['field']]==edit['old_value'];by[edit['case_id']][edit['field']]=edit['new_value']
        data=dict(revision=change['revision'],cases=list(by.values()),business_as_of=packet['as_of'])
    elif node=='S3':
        packet=load(work/'S2.json')
        data=dict(cases=[dict(decide(c),adopted_revision=packet['revision']) for c in packet['cases']],training_only=True)
    elif node=='S4':data=dict(topics=load(work/'S1.json')['rules'],decision_authority='机构/伦理委员会，非秘书',release='待机构审阅')
    elif node in ('S5','S6','S7'):
        rows=load(work/'S3.json')['cases'];name=OUTPUTS[{'S5':0,'S6':1,'S7':2}[node]]
        {'S5':book,'S6':slides,'S7':notes}[node](output/name,rows)
        data=dict(output=name,sha256=digest(output/name))
    elif node=='S8':
        for dep in ('S5','S6','S7'):
            item=load(work/(dep+'.json'));assert digest(output/item['output'])==item['sha256'],'Output modified after generation'
        data=check_office(output,load(work/'S3.json')['cases'])
    else:data=dict(gate='READY_FOR_INSTITUTIONAL_TRAINING_REVIEW',ethics_approval=False,clinical_clearance=False,system_upload_performed=False)
    artifact=node+'.json';dump(work/artifact,data)
    state['stages'][node]=dict(dependencies=DEPS[node],artifact=artifact,sha256=digest(work/artifact),completed_at_utc=datetime.datetime.now(datetime.timezone.utc).isoformat())
    dump(statepath,state)
    return state


def receipt(root,output,work,state):
    assert set(state['stages'])==set(DEPS)
    dump(output/'build_receipt.json',dict(input_sha256=state['input_sha256'],workflow_sha256=digest(__file__),output_sha256={n:digest(output/n) for n in OUTPUTS},stages=state['stages'],gate=load(work/'S9.json')['gate'],case_routes={r['case_id']:r['route'] for r in load(work/'S3.json')['cases']},actual_execution_at_utc=datetime.datetime.now(datetime.timezone.utc).isoformat(),ethics_approval=False,system_upload_performed=False,scope='Actual bounded training workflow execution, not independent domain-expert approval or a solver trajectory.'))


def verify(root,output,work):
    r=load(output/'build_receipt.json');state=load(work/'state.json')
    assert r['input_sha256']==inputs(root)==state['input_sha256'],'Stale input snapshot'
    assert r['workflow_sha256']==digest(__file__)==state['workflow_sha256'],'Stale workflow'
    assert r['stages']==state['stages'] and set(r['stages'])==set(DEPS),'Missing or stale stage ledger'
    for node,entry in r['stages'].items():
        assert entry['dependencies']==DEPS[node] and all(dep in r['stages'] for dep in DEPS[node])
        assert digest(work/entry['artifact'])==entry['sha256'],'Stale intermediate '+node
    assert r['output_sha256']=={n:digest(output/n) for n in OUTPUTS},'Stale deliverable'
    assert r['case_routes']=={x['case_id']:x['route'] for x in load(work/'S3.json')['cases']}
    check_office(output,load(work/'S3.json')['cases'])
    assert r['ethics_approval'] is False and r['system_upload_performed'] is False
    return dict(pass_all=True,gate=r['gate'],stages=len(r['stages']),scope='Version/hash and authored consistency verification, not a general semantic judge')


def main():
    p=argparse.ArgumentParser();p.add_argument('command',choices=['build','verify','stage']);p.add_argument('node',nargs='?');p.add_argument('--input-dir',type=Path,required=True);p.add_argument('--output-dir',type=Path,required=True);p.add_argument('--work-dir',type=Path,required=True);a=p.parse_args()
    root=a.input_dir.resolve();output=a.output_dir.resolve();work=a.work_dir.resolve();output.mkdir(parents=True,exist_ok=True);work.mkdir(parents=True,exist_ok=True)
    try:
        if a.command=='build':
            assert not (work/'state.json').exists(),'Choose a fresh work-dir; prior run must not be silently overwritten'
            for node in DEPS:state=stage(node,root,output,work)
            receipt(root,output,work,state);result=verify(root,output,work)
        elif a.command=='stage':
            state=stage(a.node,root,output,work);result=dict(completed=a.node)
            if a.node=='S9':receipt(root,output,work,state)
        else:result=verify(root,output,work)
        print(json.dumps(result,ensure_ascii=False))
    except Exception as exc:
        print(json.dumps(dict(pass_all=False,error=str(exc),scope='No approval or successful completion inferred'),ensure_ascii=False));raise SystemExit(2)


if __name__=='__main__':main()
