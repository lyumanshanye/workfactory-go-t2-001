"""Frozen facility evidence gate. No settings, repairs or regulator approvals."""
import argparse,csv,datetime,hashlib,json
from pathlib import Path
from openpyxl import load_workbook
def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()
def main():
    parser=argparse.ArgumentParser();parser.add_argument('--input-dir',type=Path,required=True);parser.add_argument('--out',type=Path,required=True);a=parser.parse_args()
    sk=Path(__file__).resolve().parents[1];sop=json.loads((sk/'references/team-sop.json').read_text());root=a.input_dir
    b=load_workbook(root/'S05_设备冻结清单.xlsx',data_only=True,read_only=True);rows=list(b['设备'].values);assets=[dict(zip(rows[0],r)) for r in rows[1:]];b.close()
    service=json.loads((root/'S06_受控维修证据.json').read_text());logs=list(csv.DictReader((root/'S07_历史与培训登记.csv').open()));bylog={r['asset_id']:r for r in logs};assert len(bylog)==len(logs)==len(assets)
    b=load_workbook(root/'S08_转出设备沟通登记.xlsx',data_only=True,read_only=True);rows=list(b.active.values);transfers={r[0]:dict(zip(rows[0],r)) for r in rows[1:]};b.close()
    result=[]
    for asset in assets:
        cid=asset['asset_id'];model=sop['models'].get(asset['model']);log=bylog[cid];issues=[];adopted=[];versions=[];target=sop['software_target_in_frozen_us_notice']
        state=None
        if model is None:state='OUTSIDE_GIVEN_MATRIX';issues.append('S01四型号矩阵未匹配，非其他召回均已排除')
        elif asset['udi_di']!=model['udi']:state='HOLD_IDENTITY';issues.append('型号/UDI-DI不一致或缺失，需实物身份核查')
        else:
            candidates=[r for r in service['records'] if r['asset_id']==cid and r['observed_at']==service['as_of'] and r['authority']=='primary']
            adopted=[r['evidence_id'] for r in candidates];versions=sorted({r['installed_version'] for r in candidates})
            if len(versions)>1:state='HOLD_SERVICE_CONFLICT';issues.append('同观察时点平级维修证据矛盾，不能择新消息覆盖')
            elif not versions or versions[0] in ['',None,'UNKNOWN']:issues.append('缺当前软件实证，需临床工程核查')
            elif versions[0]!=target:issues.append('版本非冻结美国通知目标，核型号/地区适用及更新；不得自行降级')
            if log['current_addendum_checked']!='yes':issues.append('最新使用手册补充说明未完成核验')
            if log['sensor_history']=='unknown':issues.append('历史雾化/传感器情况未知，需核查而非填无风险')
            elif log['sensor_history']=='deposit_concern' and log['sensor_service_resolved']!='yes':issues.append('历史传感器沉积事项尚未服务闭环，软件更新不能代替')
            if log['current_guidance_briefed']!='yes':issues.append('新版雾化指导告知/培训未确认')
            if asset['custody']=='transferred':
                t=transfers.get(cid)
                if not t or t['latest_notice_and_addendum_ack'] is not True:issues.append('转出接收方未确认新版通知及补充说明，须追踪沟通')
            if state is None:state='ACTION_PENDING' if issues else 'READY_FOR_AUTHORIZED_REVIEW'
        t=transfers.get(cid)
        result.append(dict(asset_id=cid,model=asset['model'],udi_di=asset['udi_di'],custody=asset['custody'],state=state,adopted_service_ids=adopted,observed_versions=versions,gaps=issues,transfer_recipient=t['recipient_role'] if t else None,transfer_current_ack=t['latest_notice_and_addendum_ack'] if t else None,clinical_release=False,external_send=False))
    counts={s:sum(r['state']==s for r in result) for s in sop['states']}
    payload=dict(task_id='MED-TIME-004',executed_at_utc=datetime.datetime.now(datetime.timezone.utc).isoformat(),sop_id=sop['sop_id'],input_sha256={p.name:sha(p) for p in root.iterdir() if p.is_file()},skill_sha256={str(p.relative_to(sk)):sha(p) for p in sorted(sk.rglob('*')) if p.is_file() and '__pycache__' not in p.parts},service_observation=service['as_of'],assets=result,counts=counts,clinical_release=False,external_send=False,scope='Actual deterministic internal evidence gate, fictional records; not FDA or clinical authorization, not a solver trajectory.')
    a.out.parent.mkdir(parents=True,exist_ok=True);a.out.write_text(json.dumps(payload,ensure_ascii=False,indent=2));print(json.dumps(dict(counts=counts,clinical_release=False,external_send=False),ensure_ascii=False))
if __name__=='__main__':main()
