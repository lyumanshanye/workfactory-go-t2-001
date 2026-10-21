"""Editorial trace binding, deliberately not a legal classifier."""
import argparse,datetime,hashlib,json
from pathlib import Path
from docx import Document
from openpyxl import load_workbook

def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()

def build(input_dir,handout,ledger):
    files=sorted(p for p in input_dir.iterdir() if p.is_file())
    if len(files)!=5:raise ValueError('Expected the five frozen input files')
    context=json.loads((input_dir/'S05_情境与消息.json').read_text())
    if context['review_status']!='not_reviewed':raise ValueError('No external review approval is provided')
    Document(handout)
    book=load_workbook(ledger,data_only=True)
    rows=list(book['纠错台账'].values)
    if list(rows[0])!=['段号','处理','依据','讲义定位','复核动作']:raise ValueError('Incorrect editorial columns')
    entries=[dict(zip(rows[0],r)) for r in rows[1:] if any(v is not None for v in r)]
    if len(entries)!=8 or {e['段号'] for e in entries}!={f'D{i:02}' for i in range(1,9)}:raise ValueError('Each original paragraph needs one traceable record')
    for e in entries:
        if e['处理'] not in ['已校订','待核'] or not all(isinstance(v,str) and v.strip() for v in e.values()):raise ValueError('Incomplete editorial action')
    case_rows=list(book['情境练习'].values)
    if list(case_rows[0])!=['情境','处理建议','依据','下一动作']:raise ValueError('Incorrect case columns')
    cases=[dict(zip(case_rows[0],r)) for r in case_rows[1:] if any(v is not None for v in r)]
    if len(cases)!=len(context['cases']) or {c['情境'] for c in cases}!={c['id'] for c in context['cases']}:raise ValueError('Case coverage differs from given context')
    if any(not all(isinstance(v,str) and v.strip() for v in c.values()) for c in cases):raise ValueError('Unresolved case trace is empty')
    skill=Path(__file__).resolve().parents[1]
    return dict(at_utc=datetime.datetime.now(datetime.timezone.utc).isoformat(),status='DRAFT_FOR_ARCHIVE_REVIEW',published=False,expert_approved=False,source_sha256={p.name:sha(p) for p in files},artifact_sha256={p.name:sha(p) for p in [handout,ledger]},skill_file_hashes={str(p.relative_to(skill)):sha(p) for p in [skill/'SKILL.md',skill/'scripts/build_evidence.py']},paragraphs=entries,cases=cases,scope='Editorial completeness and byte binding only; not domain correctness, release approval, or solver trajectory')

def main():
    p=argparse.ArgumentParser();p.add_argument('--input-dir',type=Path,required=True);p.add_argument('--handout',type=Path,required=True);p.add_argument('--ledger',type=Path,required=True);p.add_argument('--out',type=Path,required=True);a=p.parse_args()
    result=build(a.input_dir,a.handout,a.ledger)
    a.out.parent.mkdir(parents=True,exist_ok=True);a.out.write_text(json.dumps(result,ensure_ascii=False,indent=2)+'\n')
    print(json.dumps(dict(status=result['status'],paragraphs=len(result['paragraphs']),cases=len(result['cases']),published=False)))

if __name__=='__main__':main()
