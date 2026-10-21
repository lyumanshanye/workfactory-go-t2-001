"""Local, source-bound presentation remediation build and dependency verification.

Original reference implementation. No real release, approval or AT certification.
"""
import argparse,csv,datetime,hashlib,json,shutil,subprocess,sys,tempfile,zipfile,os
from pathlib import Path
from decimal import Decimal
from pptx import Presentation
from pptx.util import Inches,Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE,MSO_CONNECTOR
from pptx.enum.chart import XL_CHART_TYPE,XL_LEGEND_POSITION,XL_LABEL_POSITION
from pptx.chart.data import CategoryChartData
from pptx.oxml.xmlchemy import OxmlElement
from docx import Document
from docx.shared import Cm
from openpyxl import Workbook,load_workbook
from openpyxl.styles import Font,PatternFill,Alignment
from openpyxl.worksheet.pagebreak import Break
from pypdf import PdfReader
import cairosvg

NAMES=['N0_freeze','N1_authority','N2_metrics','N3_story','N4_deck','N5_structure','N6_render','N7_handoff','N8_release']
PARENTS={NAMES[0]:[],NAMES[1]:[NAMES[0]],NAMES[2]:[NAMES[0],NAMES[1]],NAMES[3]:[NAMES[1],NAMES[2]],NAMES[4]:[NAMES[3]],NAMES[5]:[NAMES[1],NAMES[2],NAMES[4]],NAMES[6]:[NAMES[4],NAMES[5]],NAMES[7]:[NAMES[1],NAMES[2],NAMES[5],NAMES[6]],NAMES[8]:[NAMES[7]]}
OUTPUTS=['repaired_deck.pptx','accessibility_review.xlsx','handoff.docx','workflow.py','release_receipt.json']
DECOR='http://schemas.microsoft.com/office/drawing/2017/decorative'
DECOR_URI='{C183D7F6-B498-43B3-948B-1728B52AA6E4}'
OFFICE=os.environ.get('DE004_OFFICE_BINARY','libreoffice')
PDFTOPPM=os.environ.get('DE004_PDFTOPPM_BINARY','pdftoppm')
PNS='http://schemas.openxmlformats.org/presentationml/2006/main'
ANS='http://schemas.openxmlformats.org/drawingml/2006/main'
NS={'p':PNS,'a':ANS,'adec':DECOR}

def sha(p):return hashlib.sha256(Path(p).read_bytes()).hexdigest()
def now():return datetime.datetime.now(datetime.timezone.utc).isoformat()
def read(p):return json.loads(Path(p).read_text())
def write(p,value):Path(p).write_text(json.dumps(value,ensure_ascii=False,indent=2))
def rows(p):
    with Path(p).open(newline='') as f:return list(csv.DictReader(f))
def hashes(folder):return {str(p.relative_to(folder)):sha(p) for p in sorted(folder.rglob('*')) if p.is_file()}
def one(s):return s._element.xpath('.//p:cNvPr')[0]
def is_decorative(s):return bool(one(s).xpath('.//*[local-name()="decorative" and (@val="1" or @val="true")]'))
def alt(s,text):one(s).set('descr',text)
def decorative(s):
    c=one(s);exts=c.find('{'+ANS+'}extLst')
    if exts is None:exts=OxmlElement('a:extLst');c.append(exts)
    ext=OxmlElement('a:ext');ext.set('uri',DECOR_URI)
    from lxml import etree
    tag=etree.SubElement(ext,'{'+DECOR+'}decorative',nsmap={'adec':DECOR});tag.set('val','1');exts.append(ext)
    c.attrib.pop('descr',None)
def color(hex):return RGBColor.from_string(hex)
def text_box(slide,name,text,left,top,width,height,size=22,bold=False,fill=None):
    s=slide.shapes.add_textbox(Inches(left),Inches(top),Inches(width),Inches(height));s.name=name
    tf=s.text_frame;tf.word_wrap=True;tf.clear();tf.margin_left=Inches(.08);tf.margin_top=Inches(.04)
    for i,line in enumerate(text.split('\n')):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.text=line;p.space_after=Pt(10)
        for run in p.runs:run.font.name='Noto Sans CJK SC';run.font.size=Pt(size);run.font.bold=bold;run.font.color.rgb=color(fill or '173C45')
    return s
def slide_start(prs,num,title,draft=False):
    sl=prs.slides.add_slide(prs.slide_layouts[5]);sl.name=f'S{num:02}'
    sl.background.fill.solid();sl.background.fill.fore_color.rgb=color('F4F7F7')
    t=sl.shapes.title;t.name='title';t.left=Inches(.65);t.top=Inches(.45);t.width=Inches(12);t.height=Inches(.8)
    t.text='服务复盘' if draft else title
    for p in t.text_frame.paragraphs:
        for r in p.runs:r.font.name='Noto Sans CJK SC';r.font.size=Pt(30);r.font.bold=True;r.font.color.rgb=color('173C45')
    if not draft:
        accent=sl.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(.65),Inches(.22),Inches(.55),Inches(.06));accent.name='accent';accent.fill.solid();accent.fill.fore_color.rgb=color('006C70');accent.line.fill.background();decorative(accent)
    return sl

def metric_table(root):
    data=sorted(rows(root/'S03_monthly_metrics.csv'),key=lambda r:int(r['month']));result=[]
    assert len(data)==12 and [int(r['month']) for r in data]==list(range(1,13)),'Need all 12 unique ordered months'
    for r in data:
        n,c,d=map(int,[r['received'],r['completed'],r['duplicate_completed']])
        assert 0<=d<=c<=n
        result.append(dict(month=int(r['month']),received=n,completed=c,duplicate_completed=d,eligible=n-d,eligible_completed=c-d))
    q=[]
    for i in range(4):
        part=result[i*3:i*3+3];n=sum(r['eligible'] for r in part);c=sum(r['eligible_completed'] for r in part)
        q.append(dict(quarter='Q'+str(i+1),received=n,completed=c,rate_pct=str((Decimal(c)/n*100).quantize(Decimal('.01')))))
    n=sum(r['eligible'] for r in result);c=sum(r['eligible_completed'] for r in result)
    return dict(months=result,quarters=q,total_received=n,total_completed=c,total_duplicate=sum(r['duplicate_completed'] for r in result),rate_pct=str((Decimal(c)/n*100).quantize(Decimal('.01'))),raw_received=sum(r['received'] for r in result),raw_completed=sum(r['completed'] for r in result),scope='原创虚构服务记录，只做内部设计练习，不是真实机构统计/因果研究')

def authority_context(root):
    """Record source event dates; never substitute acquisition/build time."""
    cfg=read(root/'S05_approved_content.json')
    assert cfg['version'] in ['R2','R3'] and cfg['authority_role']=='内容负责人' and cfg['approval_scope']=='INTERNAL_REVIEW_ONLY'
    register=rows(root/'S06_revision_register.csv');assert len({r['version'] for r in register})==len(register)
    assert {'R1','R2','M-LATE'}<={r['version'] for r in register}
    for r in register:datetime.datetime.strptime(r['date_local'],'%Y-%m-%d %H:%M')
    approved=next((r for r in register if r['version']==cfg['version']),None)
    if approved:assert approved['date_local']==cfg['approved_at_local'] and approved['role']==cfg['authority_role'],'Approval date/role conflicts with revision register'
    else:
        # Controlled R3 test updates S05 explicitly; it is not a historical S06 row.
        register.append(dict(version=cfg['version'],date_local=cfg['approved_at_local'],role=cfg['authority_role'],scope='S05新增批准；仅内部复核',status='APPROVED_FOR_INTERNAL_REVIEW_FROM_S05'))
    register.sort(key=lambda r:(r['date_local'],r['version']))
    acquisition=next(p.text for p in Document(root/'S16_official_source_packet.docx').paragraphs if '抓取于' in p.text)
    return dict(timezone='Asia/Shanghai',events=register,current_revision=cfg['version'],current_approved_at_local=cfg['approved_at_local'],business_as_of_local='2026-10-20 16:00',source_acquisition_statement=acquisition,source_time_precision='S16仅给2026-10-20/21日期范围；逐源精确UTC未知，不由当前运行时间补造',event_source='S06_revision_register.csv；当前批准另核S05_approved_content.json；无权建议另核S09_late_message.md')

def narrative(root,metrics):
    cfg=read(root/'S05_approved_content.json');assert cfg['version'] in ['R2','R3'] and cfg['authority_role']=='内容负责人'
    assert cfg['approval_scope']=='INTERNAL_REVIEW_ONLY' and cfg['data_sha256']==sha(root/'S03_monthly_metrics.csv')
    assert len(cfg['slides'])==12 and [s['id'] for s in cfg['slides']]==[f'S{i:02}' for i in range(1,13)]
    parameters={k:str(metrics[k]) for k in ['total_received','total_completed','total_duplicate','rate_pct','raw_received','raw_completed']}
    for s in cfg['slides']:
        s['body']=s['body'].format(**parameters);s['notes']=s['notes'].format(**parameters)
    cfg['authority_timeline']=authority_context(root)
    return cfg

def make_deck(root,plan,metrics,dest,draft=False):
    prs=Presentation();prs.slide_width=Inches(13.333333);prs.slide_height=Inches(7.5)
    for i,content in enumerate(plan['slides'],1):
        sl=slide_start(prs,i,content['title'],draft)
        if i==4:
            cd=CategoryChartData();cd.categories=[r['quarter'] for r in metrics['quarters']]
            cd.add_series('合格受理',[r['received'] for r in metrics['quarters']]);cd.add_series('合格完成',[r['completed'] for r in metrics['quarters']])
            frame=sl.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED,Inches(.9),Inches(1.55),Inches(11.5),Inches(4.5),cd);frame.name='quarter_chart';ch=frame.chart
            ch.has_legend=True;ch.legend.position=XL_LEGEND_POSITION.BOTTOM;ch.legend.include_in_layout=False;ch.legend.font.size=Pt(16)
            ch.value_axis.maximum_scale=600;ch.value_axis.minimum_scale=0;ch.value_axis.tick_labels.font.size=Pt(14);ch.category_axis.tick_labels.font.size=Pt(18)
            for series,hue in zip(ch.series,['006C70','A74B24']):series.format.fill.solid();series.format.fill.fore_color.rgb=color(hue);series.format.line.fill.background()
            plot=ch.plots[0];plot.has_data_labels=True;plot.data_labels.position=XL_LABEL_POSITION.OUTSIDE_END;plot.data_labels.font.size=Pt(14)
            alt(frame,plan['version']+'去重口径：'+ '；'.join(r['quarter']+'合格受理'+str(r['received'])+'、合格完成'+str(r['completed']) for r in metrics['quarters'])+'。合计'+str(metrics['total_completed'])+'/'+str(metrics['total_received'])+'，非真实机构统计。')
            text_box(sl,'body',content['body'],.9,6.15,11.7,.5,20)
        elif i==5:
            frame=sl.shapes.add_table(5,4,Inches(.9),Inches(1.65),Inches(11.5),Inches(3.8));frame.name='quarter_table';table=frame.table
            values=[['季度','合格受理','合格完成','完成率']]+[[r['quarter'],str(r['received']),str(r['completed']),r['rate_pct']+'%'] for r in metrics['quarters']]
            for ri,row in enumerate(values):
                for ci,val in enumerate(row):
                    cell=table.cell(ri,ci);cell.text=val;cell.fill.solid();cell.fill.fore_color.rgb=color('173C45' if ri==0 else 'FFFFFF')
                    for p in cell.text_frame.paragraphs:
                        for run in p.runs:run.font.name='Noto Sans CJK SC';run.font.size=Pt(22);run.font.color.rgb=color('FFFFFF' if ri==0 else '173C45')
            alt(frame,'季度明细表：每行按季度、合格受理、合格完成、完成率四列。全年按合计分子/合计分母为'+metrics['rate_pct']+'%，不是季度完成率的简单平均。')
            text_box(sl,'body',content['body'],.9,5.8,11.5,.8,20)
        elif i==6:
            group=sl.shapes.add_group_shape();group.name='service_flow'
            nodes=rows(root/'S10_workflow_nodes.csv');assert len(nodes)==6
            for j,n in enumerate(nodes):
                x=.9+(j%3)*4.1;y=1.8+(j//3)*2.05
                box=group.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,Inches(x),Inches(y),Inches(3.35),Inches(1.1));box.name=n['id'];box.text=n['label'];box.fill.solid();box.fill.fore_color.rgb=color('FFFFFF');box.line.color.rgb=color('006C70')
                for p in box.text_frame.paragraphs:
                    for run in p.runs:run.font.name='Noto Sans CJK SC';run.font.size=Pt(21);run.font.color.rgb=color('173C45')
            # Logical sequence is explicit in text and group alt, not color-only.
            for edge in rows(root/'S11_workflow_edges.csv'):
                j=int(edge['from'][1:])-1
                if j!=2:
                    x=4.35+(j%3)*4.1;y=2.1+(j//3)*2.05
                    arrow=group.shapes.add_shape(MSO_SHAPE.CHEVRON,Inches(x),Inches(y),Inches(.45),Inches(.4));arrow.name=edge['from']+'_'+edge['to'];arrow.fill.solid();arrow.fill.fore_color.rgb=color('006C70');arrow.line.fill.background();decorative(arrow)
                else:
                    for k,(x1,y1,x2,y2) in enumerate([(10.77,2.9,10.77,3.35),(10.77,3.35,2.57,3.35),(2.57,3.35,2.57,3.65)]):
                        line=group.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(x1),Inches(y1),Inches(x2),Inches(y2));line.name='F3_F4_segment'+str(k);line.line.color.rgb=color('006C70');line.line.width=Pt(2);decorative(line)
                    arrow=group.shapes.add_shape(MSO_SHAPE.DOWN_ARROW,Inches(2.4),Inches(3.54),Inches(.34),Inches(.28));arrow.name='F3_F4';arrow.fill.solid();arrow.fill.fore_color.rgb=color('006C70');arrow.line.fill.background();decorative(arrow)
            alt(group,'原创服务流程，按编号1到6阅读：'+'→'.join(n['label'] for n in nodes)+'。第3步存在疑问须转人工，不得自动视为批准。')
            text_box(sl,'body',content['body'],.9,6.15,11.5,.65,20)
        elif i==9:
            with tempfile.TemporaryDirectory(prefix='deck-illustration-') as tmp:
                png=Path(tmp)/'illustration.png';svg=(root/'S12_interface_illustration.svg').read_text().replace('font-size="21"','font-size="25"').replace('font-size="24"','font-size="25"')
                cairosvg.svg2png(bytestring=svg.encode(),write_to=str(png),output_width=1000,output_height=420)
                image=sl.shapes.add_picture(str(png),Inches(.9),Inches(1.35),width=Inches(11.4));image.name='interface_illustration'
            alt(image,'原创示意界面，不是上线截图：材料名称输入、仅本机预览按钮与待人工复核提示；尚未发送或批准。')
            text_box(sl,'body',content['body'],.9,6.3,11.5,.55,20)
        else:
            lines=content['body'].split('\n');body='\n'.join('• '+s for s in lines)
            text_box(sl,'body',body,.85,1.65,11.7,4.65,24 if len(lines)<=3 else 22)
        text_box(sl,'footer',f'原创虚构案例 · 内部待审 · {content["id"]} · {plan["version"]}',.85,7.0,11.5,.28,11)
        sl.notes_slide.notes_text_frame.text=content['notes']+f'\n来源：{content["source"]}。实际PowerPoint Checker、屏幕阅读器与外部发布审批未执行，不宣称已通过。'
        if draft:
            # Genuine old artifact defects: duplicate titles, stale notes,
            # stale numbers and bad lexical order, not a polished Gold copy.
            sl.notes_slide.notes_text_frame.text='R1：受理1620、完成1427；全部无障碍测试已通过，可直接发布。'
            for sh in sl.shapes:
                one(sh).attrib.pop('descr',None)
            tree=sl.shapes._spTree;objects=list(tree)[2:]
            for element in objects:tree.remove(element)
            for element in reversed(objects):tree.append(element)
        else:
            # Keep title first among informative objects; decorative accent last.
            tree=sl.shapes._spTree;objects=list(tree)[2:];decor=[e for e in objects if e.xpath('.//*[local-name()="decorative" and (@val="1" or @val="true")]') and e.tag!='{'+PNS+'}grpSp']
            for e in decor:tree.remove(e);tree.append(e)
    prs.save(dest)

def inspect_deck(path,metrics):
    prs=Presentation(path);assert len(prs.slides)==12
    result=[];titles=[]
    for i,sl in enumerate(prs.slides,1):
        active=[s for s in sl.shapes if not is_decorative(s)];title=sl.shapes.title
        assert title is not None and title.text.strip();titles.append(title.text.strip())
        assert active[0].name=='title',f'S{i:02} title not first'
        assert sl.name==f'S{i:02}' and len({s.shape_id for s in sl.shapes})==len(sl.shapes)
        for s in sl.shapes:
            assert 0<=s.left and 0<=s.top and s.left+s.width<=prs.slide_width+50 and s.top+s.height<=prs.slide_height+50,(i,s.name,'outside slide')
            if s.shape_type==6 or s.shape_type==13 or s.has_chart or s.has_table:
                assert len(one(s).get('descr',''))>25 and not is_decorative(s),(i,s.name,'informative object missing alt or decorative')
        assert sl.has_notes_slide and '屏幕阅读器' in sl.notes_slide.notes_text_frame.text
        assert '全部无障碍测试已通过' not in sl.notes_slide.notes_text_frame.text
        result.append(dict(slide_id=sl.name,title=title.text,active_order=[s.name for s in active],decorative=[s.name for s in sl.shapes if is_decorative(s)],notes_present=True,shape_count=len(sl.shapes),scope='OOXML lexical/navigation structure and explicit alt; not actual AT speech'))
    assert len(set(titles))==12
    chart=next(s.chart for s in prs.slides[3].shapes if s.has_chart)
    assert [int(x) for x in chart.series[0].values]==[r['received'] for r in metrics['quarters']]
    assert [int(x) for x in chart.series[1].values]==[r['completed'] for r in metrics['quarters']]
    alt_chart=one(next(s for s in prs.slides[3].shapes if s.has_chart)).get('descr','')
    assert str(metrics['total_completed'])+'/'+str(metrics['total_received']) in alt_chart
    table=next(s.table for s in prs.slides[4].shapes if s.has_table)
    assert [[table.cell(i,j).text for j in range(4)] for i in range(1,5)]==[[r['quarter'],str(r['received']),str(r['completed']),r['rate_pct']+'%'] for r in metrics['quarters']]
    return result

def office_render(deck,work):
    render=work/'render';render.mkdir(exist_ok=False)
    with tempfile.TemporaryDirectory(prefix='deck-remediation-lo-') as profile:
        p=subprocess.run([OFFICE,'-env:UserInstallation='+Path(profile).as_uri(),'--headless','--convert-to','pdf','--outdir',str(render),str(deck)],capture_output=True,text=True,timeout=120)
    assert p.returncode==0,p.stdout+p.stderr
    pdf=render/'repaired_deck.pdf';assert len(PdfReader(pdf).pages)==12
    p=subprocess.run([PDFTOPPM,'-scale-to','1280','-png',str(pdf),str(render/'slide')],capture_output=True,text=True,timeout=120);assert p.returncode==0,p.stderr
    pngs=list(render.glob('slide-*.png'));assert len(pngs)==12
    version=subprocess.run([OFFICE,'--version'],capture_output=True,text=True,timeout=30);assert version.returncode==0
    raster_version=subprocess.run([PDFTOPPM,'-v'],capture_output=True,text=True,timeout=30);assert raster_version.returncode==0
    return dict(tool='LibreOffice + pdftoppm',actual_office_executable=shutil.which(OFFICE),actual_office_version=version.stdout.strip(),actual_raster_executable=shutil.which(PDFTOPPM),actual_raster_version=(raster_version.stdout+raster_version.stderr).splitlines()[0],pdf_sha256=sha(pdf),png_sha256={p.name:sha(p) for p in pngs},pages=12,actual_time_utc=now(),manual_visual_review='PENDING_DESIGNER',native_powerpoint_checker='NOT_RUN',screen_reader='NOT_RUN',pdf_ua='NOT_CLAIMED')

def make_handoff(root,out,work,plan,metrics,audit,render):
    wb=Workbook();s=wb.active;s.title='逐页复核';s.append(['页ID','标题','原D0问题与定位','实际修复/信息顺序','当前证据定位','验证状态与未测责任'])
    vocabulary={'title':'标题','body':'正文','footer':'状态页脚','quarter_chart':'季度原生图','quarter_table':'季度原生表','service_flow':'六步流程组','interface_illustration':'示意图'}
    special={'S03':'；正文旧计数/未去重','S04':'；原生图沿用R1季度量且无alt','S05':'；原生表沿用R1且无alt','S06':'；流程分组缺语义alt','S09':'；示意图缺上下文alt'}
    for i,r in enumerate(audit):
        problem='S02/S13：重复标题、逆序对象、旧未审通过备注'+special.get(r['slide_id'],'')
        fix=' → '.join(vocabulary.get(n,n) for n in r['active_order'])+'；纯装饰'+','.join(r['decorative'])+'已标记'
        evidence=f'N5_structure.json payload.slides[{i}]；N6_render.json；render/slide-{i+1:02}.png'
        s.append([r['slide_id'],r['title'],problem,fix,evidence,'已做：OOXML/LO渲染。待设计师及有权复核者：目标PPT、真实读屏、逐页视觉。'])
    s=wb.create_sheet('口径勾稽');s.append(['季度','合格受理','合格完成','完成率公式','来源'])
    for i,r in enumerate(metrics['quarters'],2):s.append([r['quarter'],r['received'],r['completed'],f'=C{i}/B{i}','S03及S05 '+plan['version']+'去重口径'])
    s.append(['全年','=SUM(B2:B5)','=SUM(C2:C5)','=C6/B6','按合计分子/合计分母，不平均季度率'])
    for row in range(2,7):s[f'D{row}'].number_format='0.00%'
    s=wb.create_sheet('版本与边界');s.append(['事项','判断','证据'])
    for r in [['采用版本',plan['version']+'；内容负责人仅批准内部复核','S05/S06/S08，较晚商务无权变更'],['时间','业务2026-10-20 16:00上海；源抓取/程序运行UTC分别保存','不是微软功能引入日期或历史应用测试记录'],['数据性质','所有服务记录与角色均原创虚构','不作真实机构统计/政策或因果推断'],['权限','结构与LO渲染不能替代原生PowerPoint/真实读屏及发布批准','S16真实文档；本题授权边界'],['汇总',str(metrics['total_completed'])+'/'+str(metrics['total_received'])+'='+metrics['rate_pct']+'%；剔除重复'+str(metrics['total_duplicate']),'S03原字段、S04字典、S05规则'],['工作流','9个真实依赖节点；缺前置/被改产物拒绝旧回执','WORK中的节点/指标/渲染；不是可信电子签名']]:s.append(r)
    timeline=authority_context(root);started=read(node_path(work,NAMES[0]))['created_at_utc']
    for event in timeline['events']:s.append(['事件 '+event['version'],event['date_local']+' Asia/Shanghai｜'+event['role']+'｜'+event['scope']+'｜'+event['status'],'S06原始登记；R3等新批准另以S05为据，不按消息较晚决定权限'])
    s.append(['源取得日期',timeline['source_time_precision'],'S16原来源说明；未知发布日期/精确取得UTC不补造'])
    s.append(['本次实际运行','开始UTC '+started+'\n实际渲染UTC '+render['actual_time_utc'],'N0_freeze.created_at_utc/N6_render；不是业务批准或源取得时间'])
    for s in wb:
        s.freeze_panes='A2';s.auto_filter.ref=s.dimensions;s.sheet_view.showGridLines=False
        for c in s[1]:c.font=Font(bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='173C45')
        widths=[9,28,48,40,36,42] if s.title=='逐页复核' else [20,35,35,25,48] if s.title=='口径勾稽' else [20,80,70]
        for j,width in enumerate(widths,1):s.column_dimensions[__import__('openpyxl').utils.get_column_letter(j)].width=width
        for row in s:
            s.row_dimensions[row[0].row].height=72 if s.title=='逐页复核' else 52 if s.title=='版本与边界' else 32
            for c in row:c.alignment=Alignment(vertical='center',wrap_text=True)
        s.page_setup.orientation='landscape';s.page_setup.paperSize=s.PAPERSIZE_A3;s.page_setup.fitToWidth=1;s.page_setup.fitToHeight=1;s.sheet_properties.pageSetUpPr.fitToPage=True;s.print_title_rows='1:1'
        if s.title=='逐页复核':s.page_setup.fitToHeight=0;s.row_breaks.append(Break(id=7))
        if s.title=='口径勾稽':
            for row in range(2,7):
                for col in ['B','C','D']:s[f'{col}{row}'].alignment=Alignment(horizontal='center',vertical='center',wrap_text=True)
    raw=work/'accessibility_review.xlsx';wb.save(raw)
    with tempfile.TemporaryDirectory(prefix='deck-review-calc-') as profile:
        p=subprocess.run([OFFICE,'-env:UserInstallation='+Path(profile).as_uri(),'--headless','--convert-to','xlsx','--outdir',str(out),str(raw)],capture_output=True,text=True,timeout=90)
    assert p.returncode==0 and (out/'accessibility_review.xlsx').exists(),p.stdout+p.stderr
    v=load_workbook(out/'accessibility_review.xlsx',data_only=True)['口径勾稽'];assert v['B6'].value==metrics['total_received'] and v['C6'].value==metrics['total_completed']
    d=Document();sec=d.sections[0];sec.top_margin=sec.bottom_margin=Cm(1.8);d.styles['Normal'].font.name='Noto Sans CJK SC';d.styles['Normal'].font.size=Pt(10.5)
    d.add_heading('演示修复交接｜内部待审',0)
    sections=[('结论与数据',f'采用{plan["version"]}。合格受理{metrics["total_received"]}、合格完成{metrics["total_completed"]}、完成率{metrics["rate_pct"]}%，剔除重复完成记录{metrics["total_duplicate"]}。原始{metrics["raw_completed"]}/{metrics["raw_received"]}只作R1对照。指标文字、季度原生图表与表格、alt及讲者备注同源。全部业务为原创虚构，不是机构统计或因果效果。'),('原生修复与可编辑性','12页独特标题与真实标题占位符；信息对象按标题→正文/图表/表格/分组→状态页脚顺序，纯装饰条使用真实decorative扩展。季度图为原生chart及嵌入数据，季度表原生table，六步图为可编辑group；示意图有上下文alt，文字对象不重复堆alt。背景不盖正文，不用白字/移出画布隐藏重要标题。'),('权限与时间','R1草稿不再是当前口径；S05内容负责人仅有内部内容修订权限，S09更晚商务建议无权恢复原数字或把状态改为已发布。业务基准2026-10-20 16:00 Asia/Shanghai；真实网页抓取与运行UTC另存，未知官方发布日期不补造。角色和hash仅是题内材料/字节绑定，不是可认证的真实签章。'),('实际做过与没有做过','实际执行OOXML结构、原生数据勾稽、LibreOffice的12页PDF与PNG导出，详细见WORK/render和release_receipt。没有执行Windows/Mac PowerPoint Accessibility Checker、真实屏幕阅读器、PDF/UA认证、真实用户研究或外部发布批准。不得把缺失测试填写通过。结构/几何检查不证明所有视觉排版和AT体验正确。'),('依赖与复现','workflow.py build --input-dir INPUT --output-dir NEW_OUTPUT --work-dir NEW_WORK；verify使用同参数只读检查。stage --name N0_freeze等按S19依赖执行，缺前置即失败，WORK/输入/输出必须分开且不嵌套。输出只保留指定5文件，PDF/PNG/节点放WORK，不把运行中间件冒最终交付。需python-pptx、openpyxl、python-docx、pypdf、cairosvg、LibreOffice Impress/Calc和pdftoppm。'),('变更与下一责任人','数据负责人核对去重字段与输入哈希，设计师查12页视觉/叠放/替代描述，使用目标PowerPoint和屏幕阅读器逐页验证再交有权负责人决定发布。合法R3更新须绑定新数据并新目录全量重建，图/表/alt/备注/工作簿与新回执一起更新；旧回执必须失效，不可只改文字。缺依赖/工具时报错不产发布通过。'),('专业依据与局限','S16微软阅读顺序/alt/形状树/非视觉属性文档；S17固定Open XML SDK v3.0.1描述和S18 MIT许可。spTree词法顺序可证层叠与导航，不单凭它宣称所有PowerPoint版本的真实读屏结果；decorative来自官方类型/SDK子元素与实际PPT标记的结构对照。示意图原创，不加载外部资源，不含真实个人信息。')]
    for title,body in sections:
        if title=='实际做过与没有做过':d.add_page_break()
        d.add_heading(title,1);d.add_paragraph(body)
        if title=='权限与时间':
            table=d.add_table(rows=1,cols=3);table.style='Table Grid'
            for cell,text in zip(table.rows[0].cells,['事件/依据','原日期（Asia/Shanghai）','角色、范围与采用处理']):cell.text=text
            for event in timeline['events']:
                labels={'DRAFT':'未审批旧稿','APPROVED_FOR_INTERNAL_REVIEW':'仅内部复核批准','APPROVED_FOR_INTERNAL_REVIEW_FROM_S05':'S05新增内部复核批准','UNAUTHORIZED':'无权覆盖/外发'}
                cells=table.add_row().cells;cells[0].text=event['version']+'｜S06' if event['version'] in ['R1','R2','M-LATE'] else event['version']+'｜S05新增批准';cells[1].text=event['date_local'];cells[2].text=event['role']+'；'+event['scope']+'；'+labels.get(event['status'],event['status'])
            d.add_paragraph('业务复核基准：'+timeline['business_as_of_local']+' Asia/Shanghai；它不是任一文件签发/批准日。源取得：'+timeline['source_time_precision']+'。未知专业文档发布日期/起施日仍未知。')
            d.add_paragraph('本次程序开始UTC：'+started+'；实际渲染UTC：'+render['actual_time_utc']+'。这是作者真实执行时点，和上表原始修订事件、业务复核基准及S16来源取得分别记录，不伪装成历史现场测试。')
    d.save(out/'handoff.docx')
    return dict(review_sha256=sha(out/'accessibility_review.xlsx'),handoff_sha256=sha(out/'handoff.docx'),authority_timeline=timeline,manual_checks_pending=True)

def validate_inputs(root):
    contract=read(root/'S19_dependency_contract.json')
    for file,h in contract['official_input_sha256'].items():assert sha(root/file)==h,'Official source drift; manual re-freeze required: '+file
    assert [r['id'] for r in rows(root/'S10_workflow_nodes.csv')]==['F1','F2','F3','F4','F5','F6']
    edges=rows(root/'S11_workflow_edges.csv');assert [(e['from'],e['to']) for e in edges]==[('F1','F2'),('F2','F3'),('F3','F4'),('F4','F5'),('F5','F6')]
    return contract

def node_path(work,name):return work/(name+'.json')
def dependencies(work,name):
    result={}
    for parent in PARENTS[name]:
        p=node_path(work,parent);assert p.is_file(),'Missing dependency '+parent
        result[parent]=sha(p)
    return result

def verify_prefix(out,work,next_name):
    index=NAMES.index(next_name)
    for previous in NAMES[:index]:
        node=read(node_path(work,previous))
        assert node['name']==previous and node['parents']==dependencies(work,previous),'Prior dependency drift: '+previous
    if index:
        assert sha(out/'workflow.py')==read(node_path(work,NAMES[0]))['payload']['executor_sha256'],'Executor copy changed'
    if index>4:
        assert sha(out/'repaired_deck.pptx')==read(node_path(work,NAMES[4]))['payload']['deck_sha256'],'Deck changed after generation'
    if index>6:
        render=read(node_path(work,NAMES[6]))['payload']
        assert sha(work/'render/repaired_deck.pdf')==render['pdf_sha256'],'Rendered PDF changed'
        for file,h in render['png_sha256'].items():assert sha(work/'render'/file)==h,'Rendered PNG changed'
    if index>7:
        payload=read(node_path(work,NAMES[7]))['payload']
        assert sha(out/'accessibility_review.xlsx')==payload['review_sha256']
        assert sha(out/'handoff.docx')==payload['handoff_sha256']

def stage(root,out,work,name):
    assert name in NAMES
    assert not node_path(work,name).exists(),'Node already exists; use a fresh work directory'
    parents=dependencies(work,name);verify_prefix(out,work,name);current=hashes(root)
    if name!=NAMES[0]:assert read(node_path(work,NAMES[0]))['payload']['input_sha256']==current,'Frozen inputs changed; fresh rebuild required'
    if name==NAMES[0]:
        validate_inputs(root);payload=dict(input_sha256=current,business_as_of='2026-10-20T16:00:00+08:00',executor_sha256=sha(__file__))
    elif name==NAMES[1]:
        cfg=read(root/'S05_approved_content.json');assert cfg['authority_role']=='内容负责人' and cfg['approval_scope']=='INTERNAL_REVIEW_ONLY' and cfg['version'] in ['R2','R3'];assert cfg['data_sha256']==sha(root/'S03_monthly_metrics.csv');cfg['authority_timeline']=authority_context(root);payload=cfg
    elif name==NAMES[2]:payload=metric_table(root)
    elif name==NAMES[3]:payload=narrative(root,read(node_path(work,NAMES[2]))['payload'])
    elif name==NAMES[4]:
        make_deck(root,read(node_path(work,NAMES[3]))['payload'],read(node_path(work,NAMES[2]))['payload'],out/'repaired_deck.pptx');payload=dict(deck_sha256=sha(out/'repaired_deck.pptx'))
    elif name==NAMES[5]:payload=dict(slides=inspect_deck(out/'repaired_deck.pptx',read(node_path(work,NAMES[2]))['payload']),deck_sha256=sha(out/'repaired_deck.pptx'))
    elif name==NAMES[6]:
        assert sha(out/'repaired_deck.pptx')==read(node_path(work,NAMES[5]))['payload']['deck_sha256'],'Deck changed after structure audit'
        payload=office_render(out/'repaired_deck.pptx',work)
    elif name==NAMES[7]:payload=make_handoff(root,out,work,read(node_path(work,NAMES[3]))['payload'],read(node_path(work,NAMES[2]))['payload'],read(node_path(work,NAMES[5]))['payload']['slides'],read(node_path(work,NAMES[6]))['payload'])
    else:
        assert {p.name for p in out.iterdir() if p.is_file()}==set(OUTPUTS)-{'release_receipt.json'}
        payload=dict(output_sha256=hashes(out),release_state='INTERNAL_REVIEW_REQUIRED',external_release=False)
    write(node_path(work,name),dict(name=name,created_at_utc=now(),parents=parents,payload=payload))
    if name==NAMES[-1]:
        write(out/'release_receipt.json',dict(actual_completed_at_utc=now(),business_as_of='2026-10-20T16:00:00+08:00',revision=read(node_path(work,NAMES[1]))['payload']['version'],authority_timeline=authority_context(root),input_sha256=current,node_sha256={n:sha(node_path(work,n)) for n in NAMES},node_records={n:read(node_path(work,n)) for n in NAMES},output_sha256=hashes(out),release_state='INTERNAL_REVIEW_REQUIRED',external_release=False,manual_checks={'native_PowerPoint_Checker':'NOT_RUN','real_screen_reader':'NOT_RUN','independent_designer':'PENDING'},metrics=read(node_path(work,NAMES[2]))['payload'],scope='Actual local deterministic build/structure/LO rendering with inline node records; raw PDF/PNG retained in WORK. Not identity attestation, PowerPoint AT certification or actual publication'))

def verify(root,out,work):
    receipt=read(out/'release_receipt.json');assert receipt['input_sha256']==hashes(root)
    assert receipt['output_sha256']=={p.name:sha(p) for p in out.iterdir() if p.is_file() and p.name!='release_receipt.json'}
    assert {p.name for p in out.iterdir() if p.is_file()}==set(OUTPUTS)
    for name in NAMES:
        p=node_path(work,name);assert sha(p)==receipt['node_sha256'][name];node=read(p);assert node['name']==name and node['parents']==dependencies(work,name)
        assert receipt['node_records'][name]==node
        assert hashlib.sha256(json.dumps(node,ensure_ascii=False,indent=2).encode()).hexdigest()==receipt['node_sha256'][name]
        for parent in PARENTS[name]:assert read(node_path(work,parent))['created_at_utc']<=node['created_at_utc']
    assert read(node_path(work,NAMES[0]))['payload']['input_sha256']==hashes(root)
    assert sha(out/'workflow.py')==read(node_path(work,NAMES[0]))['payload']['executor_sha256']
    expected_timeline=authority_context(root)
    assert receipt['authority_timeline']==expected_timeline,'Current source event timeline differs'
    for node in [NAMES[1],NAMES[3],NAMES[7]]:assert read(node_path(work,node))['payload']['authority_timeline']==expected_timeline,'Timeline drift: '+node
    assert not receipt['external_release'] and all(x in ['NOT_RUN','PENDING'] for x in receipt['manual_checks'].values())
    m=metric_table(root);assert m==receipt['metrics'];inspect_deck(out/'repaired_deck.pptx',m)
    v=load_workbook(out/'accessibility_review.xlsx',data_only=True)['口径勾稽'];assert v['B6'].value==m['total_received'] and v['C6'].value==m['total_completed']
    render=read(node_path(work,NAMES[6]))['payload'];assert sha(work/'render/repaired_deck.pdf')==render['pdf_sha256']
    for name,h in render['png_sha256'].items():assert sha(work/'render'/name)==h
    assert read(node_path(work,NAMES[4]))['payload']['deck_sha256']==sha(out/'repaired_deck.pptx')
    assert read(node_path(work,NAMES[7]))['payload']['review_sha256']==sha(out/'accessibility_review.xlsx')
    assert read(node_path(work,NAMES[7]))['payload']['handoff_sha256']==sha(out/'handoff.docx')
    return dict(status='CURRENT_LOCAL_BUILD_VERIFIED',metrics=m,external_release=False,actual_verified_at_utc=now())

def main():
    parser=argparse.ArgumentParser();parser.add_argument('command',choices=['build','stage','verify']);parser.add_argument('--name',choices=NAMES);parser.add_argument('--input-dir',required=True);parser.add_argument('--output-dir',required=True);parser.add_argument('--work-dir',required=True);a=parser.parse_args()
    root,out,work=[Path(x).resolve() for x in [a.input_dir,a.output_dir,a.work_dir]]
    assert root.is_dir() and len({root,out,work})==3
    for p in [root,out,work]:
        for q in [root,out,work]:
            if p!=q:assert p not in q.parents,'Input/output/work may not nest'
    if a.command=='verify':print(json.dumps(verify(root,out,work),ensure_ascii=False));return
    if a.command=='build':
        assert not out.exists() or not any(out.iterdir());assert not work.exists() or not any(work.iterdir())
        out.mkdir(parents=True,exist_ok=True);work.mkdir(parents=True,exist_ok=True);shutil.copy2(Path(__file__).resolve(),out/'workflow.py')
        for name in NAMES:stage(root,out,work,name)
        print(json.dumps(verify(root,out,work),ensure_ascii=False))
    else:
        assert a.name
        out.mkdir(parents=True,exist_ok=True);work.mkdir(parents=True,exist_ok=True)
        if a.name==NAMES[0] and not (out/'workflow.py').exists():shutil.copy2(Path(__file__).resolve(),out/'workflow.py')
        stage(root,out,work,a.name);print(json.dumps(dict(completed=a.name,external_release=False)))

if __name__=='__main__':
    try:main()
    except Exception as exc:print(type(exc).__name__+': '+str(exc),file=sys.stderr);raise SystemExit(2)
