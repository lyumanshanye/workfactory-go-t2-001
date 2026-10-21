"""Reproduce a bounded NHANES 2017-2018 systolic-pressure analysis offline."""

import argparse
import hashlib
import importlib.metadata
import json
from datetime import datetime, timezone
from pathlib import Path

import pandas as pd
import pyreadstat
from samplics.estimation import TaylorEstimator
from samplics.utils import PopParam


def load_data(folder):
    demo, _ = pyreadstat.read_xport(folder / "DEMO_J.xpt")
    bp, _ = pyreadstat.read_xport(folder / "BPX_J.xpt")
    if not demo.SEQN.is_unique or not bp.SEQN.is_unique:
        raise ValueError("SEQN must be unique in each input")
    if not set(bp.SEQN) <= set(demo.SEQN):
        raise ValueError("blood-pressure ids absent from demographics")
    if demo.WTMEC2YR.isna().any() or not demo.WTMEC2YR.ge(0).all():
        raise ValueError("invalid exam weight")
    if not demo.WTMEC2YR.gt(0).equals(demo.RIDSTATR.eq(2)):
        raise ValueError("exam status and positive weights disagree; inspect XPT decoding")
    frame = demo.merge(bp[["SEQN", "BPXSY1", "BPXSY2", "BPXSY3", "BPXSY4"]],
                       on="SEQN", how="left", validate="one_to_one")
    readings = frame[["BPXSY1", "BPXSY2", "BPXSY3"]]
    frame["sbp"] = readings.where(readings > 0).mean(axis=1)
    base = frame.loc[frame.WTMEC2YR.gt(0)].copy()
    if not base.RIDSTATR.eq(2).all():
        raise ValueError("positive exam weights disagree with exam status; inspect XPT decoding")
    if base[["SDMVSTRA", "SDMVPSU", "RIDAGEYR", "RIAGENDR"]].isna().any().any():
        raise ValueError("missing design or domain variable")
    return frame, base


def estimate(base):
    adult = base.RIDAGEYR.ge(20)
    eligible = adult & base.sbp.notna()
    groups = {
        "all-20plus": adult,
        "male-20plus": adult & base.RIAGENDR.eq(1),
        "female-20plus": adult & base.RIAGENDR.eq(2),
        "age-20-39": base.RIDAGEYR.between(20, 39),
        "age-40-59": base.RIDAGEYR.between(40, 59),
        "age-60plus": base.RIDAGEYR.ge(60),
    }
    rows = []
    for name, members in groups.items():
        included = members & eligible
        if not included.any():
            raise ValueError("empty analysis domain: " + name)
        estimator = TaylorEstimator(PopParam.mean, alpha=0.05)
        # Zero is only a computational placeholder outside the observed domain.
        # Domain membership excludes every missing outcome from its denominator.
        estimator.estimate(y=base.sbp.fillna(0).to_numpy(), samp_weight=base.WTMEC2YR.to_numpy(),
                           stratum=base.SDMVSTRA.to_numpy(), psu=base.SDMVPSU.to_numpy(),
                           domain=included.astype(int).to_numpy())
        rows.append({"domain": name, "n_observed": int(included.sum()),
                     "n_adult_domain": int(members.sum()),
                     "n_missing": int((members & ~eligible).sum()),
                     "weight_sum": float(base.loc[included, "WTMEC2YR"].sum()),
                     "mean_mmHg": float(estimator.point_est[1]),
                     "se_mmHg": float(estimator.stderror[1]),
                     "lower95_mmHg": float(estimator.lower_ci[1]),
                     "upper95_mmHg": float(estimator.upper_ci[1]),
                     "design_df": int(estimator.degree_of_freedom)})
    return rows


def analyze(input_dir, output_dir):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    frame, base = load_data(input_dir)
    rows = estimate(base)
    result = {
        "cycle": "2017-2018", "protocol": "BPX_J traditional measurements",
        "weight": "WTMEC2YR", "readings": ["BPXSY1", "BPXSY2", "BPXSY3"],
        "estimand": "crude mean systolic mmHg among adults aged20+ with >=1 positive reading in first3",
        "variance": "Taylor linearization, with-replacement strata/PSU design; 95% t intervals",
        "missing_policy": "no imputation or item-nonresponse adjustment; report observed-case domains",
        "data_counts": {"demographics": len(frame), "mec_positive_weight": len(base),
                        "strata": int(base.SDMVSTRA.nunique()),
                        "nested_psus": len(base[["SDMVSTRA", "SDMVPSU"]].drop_duplicates())},
        "source_sha256": {name: hashlib.sha256((input_dir / name).read_bytes()).hexdigest()
                          for name in ("DEMO_J.xpt", "BPX_J.xpt")},
        "packages": {name: importlib.metadata.version(name) for name in ("samplics", "pandas", "numpy", "scipy", "pyreadstat")},
        "rows": rows,
    }
    output_dir.mkdir(parents=True, exist_ok=True)
    (output_dir / "run.json").write_text(json.dumps(result, ensure_ascii=False, indent=2, allow_nan=False) + "\n", encoding="utf-8")
    book = Workbook()
    book.properties.created = book.properties.modified = datetime(2000, 1, 1, tzinfo=timezone.utc)
    sheet = book.active
    sheet.title = "Estimates"
    headers = list(rows[0])
    sheet.append(headers)
    for row in rows:
        sheet.append([row[key] for key in headers])
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="176B70")
    for i, key in enumerate(headers, 1):
        sheet.column_dimensions[get_column_letter(i)].width = 24 if i == 1 else 19
        for cell in list(sheet.columns)[i - 1][1:]:
            if key.endswith("mmHg"):
                cell.number_format = "0.0000"
            cell.alignment = Alignment(vertical="center")
    sheet.freeze_panes = "B2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
    sheet.page_setup.fitToWidth = sheet.page_setup.fitToHeight = 1
    sheet.print_options.horizontalCentered = True
    sheet.print_area = sheet.dimensions
    sheet.print_title_rows = "1:1"
    notes = book.create_sheet("Scope")
    for key in ("cycle", "protocol", "weight", "estimand", "variance", "missing_policy"):
        notes.append([key, result[key]])
    notes.column_dimensions["A"].width = 24
    notes.column_dimensions["B"].width = 110
    for row in notes:
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="center")
        notes.row_dimensions[row[0].row].height = 32
    book.save(output_dir / "estimates.xlsx")
    return result


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--input-dir", type=Path, required=True)
    parser.add_argument("--output-dir", type=Path, required=True)
    arguments = parser.parse_args()
    report = analyze(arguments.input_dir, arguments.output_dir)
    print(json.dumps({"domains": len(report["rows"]), "counts": report["data_counts"]}))
