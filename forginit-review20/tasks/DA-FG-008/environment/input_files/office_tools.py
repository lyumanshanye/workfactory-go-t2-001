"""Office artifacts with explicit styles and printable table dimensions."""

import math
import unicodedata
import os
import tempfile
import zipfile
from datetime import datetime, timezone


def canonicalize_office(path):
    """Keep unchanged Office content byte-identical across rebuilds."""
    fd, temporary = tempfile.mkstemp(prefix=".office-", dir=path.parent)
    os.close(fd)
    try:
        with zipfile.ZipFile(path) as source, zipfile.ZipFile(temporary, "w", compression=zipfile.ZIP_DEFLATED) as target:
            for name in sorted(source.namelist()):
                entry = zipfile.ZipInfo(name, (2000, 1, 1, 0, 0, 0))
                entry.compress_type = zipfile.ZIP_DEFLATED
                entry.external_attr = source.getinfo(name).external_attr
                target.writestr(entry, source.read(name))
        os.replace(temporary, path)
        path.chmod(0o644)
    finally:
        if os.path.exists(temporary):
            os.unlink(temporary)


def text_width(value):
    return sum(2 if unicodedata.east_asian_width(c) in "WF" else 1 for c in str(value))


def workbook(path, sheets, cell_indent=0, rows_per_page=6, number_formats=None, sheet_rows_per_page=None):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter, column_index_from_string
    from openpyxl.worksheet.pagebreak import Break
    if rows_per_page is not None and (type(rows_per_page) is not int or rows_per_page < 1):
        raise ValueError("rows_per_page must be a positive integer or None")
    sheet_rows_per_page = {} if sheet_rows_per_page is None else sheet_rows_per_page
    if (not isinstance(sheet_rows_per_page, dict) or set(sheet_rows_per_page) - {s[0] for s in sheets}
            or any(value is not None and (type(value) is not int or value < 1)
                   for value in sheet_rows_per_page.values())):
        raise ValueError("sheet pagination must refer to declared sheets with positive integers or None")
    number_formats = {} if number_formats is None else number_formats
    if not isinstance(number_formats, dict) or set(number_formats) - {s[0] for s in sheets}:
        raise ValueError("number formats must refer to declared sheets")
    for name, headers, _, _ in sheets:
        columns = number_formats.get(name, {})
        if not isinstance(columns, dict):
            raise ValueError("column number formats must be a mapping")
        for column, pattern in columns.items():
            if (not isinstance(column, str) or not column.isascii() or not column.isalpha()
                    or column != column.upper() or column_index_from_string(column) > len(headers)
                    or not isinstance(pattern, str) or not pattern):
                raise ValueError("invalid column number format")
    book = Workbook()
    book.properties.created = book.properties.modified = datetime(2000, 1, 1, tzinfo=timezone.utc)
    book.remove(book.active)
    for name, headers, rows, widths in sheets:
        sheet = book.create_sheet(name)
        sheet.append(headers)
        for row in rows:
            sheet.append(row)
        for column, pattern in number_formats.get(name, {}).items():
            for cells in sheet.iter_rows(min_row=2, min_col=column_index_from_string(column),
                                         max_col=column_index_from_string(column)):
                cells[0].number_format = pattern
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for i, width in enumerate(widths, 1):
            sheet.column_dimensions[get_column_letter(i)].width = width
        for row in sheet:
            lines = 1
            for cell in row:
                cell.font = Font(name="Noto Sans CJK SC", size=10, bold=cell.row == 1,
                                 color="FFFFFF" if cell.row == 1 else "20252B")
                cell.alignment = Alignment(vertical="top", wrap_text=True, indent=cell_indent,
                    horizontal=("center" if isinstance(cell.value, (int, float)) else "left") if cell_indent else None)
                if cell.row == 1:
                    cell.fill = PatternFill("solid", fgColor="255B53")
                elif cell.row % 2 == 0:
                    cell.fill = PatternFill("solid", fgColor="EDF3F1")
                lines = max(lines, sum(max(1, math.ceil(text_width(line) / (widths[cell.column - 1] - 2)))
                                       for line in str(cell.value or "").splitlines()))
            sheet.row_dimensions[row[0].row].height = min(409, max(28, lines * 18 + 14))
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_setup.orientation = "landscape"
        sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
        sheet.print_title_rows = "1:1"
        sheet.print_options.horizontalCentered = True
        pagination = sheet_rows_per_page.get(name, rows_per_page)
        if pagination is not None and len(rows) > max(8, pagination):
            for boundary in range(pagination + 1, len(rows), pagination):
                sheet.row_breaks.append(Break(id=boundary))
        sheet.oddFooter.center.text = "&A | &P / &N"
        sheet.sheet_view.zoomScale = 85
    path.parent.mkdir(parents=True, exist_ok=True)
    book.save(path)
    # openpyxl resets modified on save; normalize that field as well as ZIP times.
    with zipfile.ZipFile(path) as archive:
        contents = {name: archive.read(name) for name in archive.namelist()}
    from xml.etree import ElementTree as ET
    core = ET.fromstring(contents["docProps/core.xml"])
    for field in ("created", "modified"):
        element = core.find("{http://purl.org/dc/terms/}" + field)
        if element is not None:
            element.text = "2000-01-01T00:00:00Z"
    contents["docProps/core.xml"] = ET.tostring(core, encoding="utf-8")
    with zipfile.ZipFile(path, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        for name, content in contents.items():
            archive.writestr(name, content)
    canonicalize_office(path)


def document(path, title, sections, subtitle=None):
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor
    doc = Document()
    doc.core_properties.created = doc.core_properties.modified = datetime(2000, 1, 1, tzinfo=timezone.utc)
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(1.8)
    section.left_margin = section.right_margin = Cm(1.9)
    normal = doc.styles["Normal"]
    normal.font.name = "Noto Sans CJK SC"
    normal.font.size = Pt(10)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Noto Sans CJK SC")
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.1
    normal.paragraph_format.keep_together = True
    for name in ("Title", "Heading 1", "Heading 2"):
        style = doc.styles[name]
        style.font.name = "Noto Sans CJK SC"
        style.element.rPr.rFonts.set(qn("w:eastAsia"), "Noto Sans CJK SC")
        style.font.color.rgb = RGBColor.from_string("255B53")
    doc.styles["Title"].font.size = Pt(20)
    doc.styles["Heading 1"].font.size = Pt(13)
    doc.add_heading(title, 0)
    if subtitle:
        doc.add_paragraph(subtitle)
    for heading, paragraphs in sections:
        doc.add_heading(heading, 1)
        for paragraph in paragraphs:
            doc.add_paragraph(paragraph)
    footer = section.footer.paragraphs[0]
    footer.add_run("第 ")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    footer.add_run(" 页")
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)
    canonicalize_office(path)
