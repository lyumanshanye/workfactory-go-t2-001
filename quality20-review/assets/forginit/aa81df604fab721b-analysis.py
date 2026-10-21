"""Reference implementation for the original retail-ledger commission."""

import argparse
from collections import Counter, defaultdict
import csv
from datetime import datetime
from decimal import Decimal
import hashlib
import json
from pathlib import Path
import sys

sys.dont_write_bytecode = True


def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def run(inputs, output, skills):
    from openpyxl import load_workbook
    sys.path.insert(0, str(inputs))
    import office_tools

    policy_path = skills / "retail-ledger-close" / "policy.json"
    policy = json.loads(policy_path.read_text())
    events = json.loads((inputs / "approval-events.json").read_text())["events"]
    cutoff = datetime.fromisoformat(policy["cutoff"])
    approved = [e for e in events if e["role"] == policy["approver_role"]
                and datetime.fromisoformat(e["occurred_at"]) <= cutoff
                and datetime.fromisoformat(e["received_at"]) <= cutoff]
    version = max(approved, key=lambda e: e["received_at"])["version"]
    config = policy["versions"][version]
    start, end = map(datetime.fromisoformat, (config["start"], config["end_exclusive"]))
    source = inputs / "Online Retail.xlsx"
    book = load_workbook(source, read_only=True, data_only=True)
    sheet = book["Online Retail"]
    data = sheet.iter_rows(values_only=True)
    headers = next(data)
    if list(headers) != ["InvoiceNo", "StockCode", "Description", "Quantity", "InvoiceDate", "UnitPrice", "CustomerID", "Country"]:
        raise ValueError("unexpected original schema")
    totals = Counter()
    samples = defaultdict(list)
    months = defaultdict(lambda: dict(sales=Decimal(0), credits=Decimal(0), missing=Decimal(0),
                                     sale_rows=0, credit_rows=0, invoices=set(), credit_invoices=set(), customers=set()))
    seen, duplicates, missing_customers = set(), 0, 0
    minimum, maximum = None, None
    raw_rows = 0
    for excel_row, row in enumerate(data, 2):
        invoice, stock, description, qty, date, price, customer, country = row
        raw_rows += 1
        signature = tuple(row)
        duplicates += signature in seen
        seen.add(signature)
        missing_customers += customer is None
        minimum = date if minimum is None else min(minimum, date)
        maximum = date if maximum is None else max(maximum, date)
        invoice = str(invoice)
        amount = Decimal(str(qty)) * Decimal(str(price))
        cancel = invoice.upper().startswith("C")
        if not start <= date < end:
            category = "outside_window"
        elif country != config["country"]:
            category = "outside_country"
        elif price <= 0:
            category = "nonpositive_price"
        elif qty == 0:
            category = "zero_quantity"
        elif cancel and qty > 0:
            category = "cancellation_sign_conflict"
        elif not cancel and qty < 0:
            category = "negative_without_c"
        elif config["require_customer"] and customer is None:
            category = "missing_customer_excluded"
        elif cancel and config["drop_cancellations"]:
            category = "cancellation_excluded"
        else:
            category = "credit" if cancel else "sale"
            month = months[date.strftime("%Y-%m")]
            if category == "credit":
                month["credits"] += -amount
                month["credit_rows"] += 1
                month["credit_invoices"].add(invoice)
            else:
                month["sales"] += amount
                month["sale_rows"] += 1
                month["invoices"].add(invoice)
                if customer is None:
                    month["missing"] += amount
                else:
                    month["customers"].add(str(customer))
        totals[category] += 1
        if len(samples[category]) < 3:
            samples[category].append(dict(excel_row=excel_row, invoice=invoice, stock=str(stock), quantity=qty,
                                         unit_price=str(price), date=date.isoformat(), customer_missing=customer is None,
                                         country=country))
    book.close()
    output.mkdir(parents=True, exist_ok=True)
    columns = ["month", "sale_rows", "credit_rows", "sale_invoices", "credit_invoices", "known_sale_customers",
               "sale_gbp", "credit_gbp", "net_recorded_gbp", "period_credit_share", "gross_invoice_mean_gbp",
               "missing_customer_sale_gbp", "missing_customer_sale_share"]
    rows = []
    for key, value in sorted(months.items()):
        sales = value["sales"]
        rows.append([key, value["sale_rows"], value["credit_rows"], len(value["invoices"]), len(value["credit_invoices"]),
                     len(value["customers"]), str(sales), str(value["credits"]), str(sales-value["credits"]),
                     str(value["credits"] / sales) if sales else "",
                     str(sales / len(value["invoices"])) if value["invoices"] else "",
                     str(value["missing"]), str(value["missing"] / sales) if sales else ""])
    with (output / "monthly.csv").open("w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle)
        writer.writerow(columns)
        writer.writerows(rows)
    audit = dict(source_sha256=sha(source), policy_sha256=sha(policy_path), approved_version=version,
                 cutoff=policy["cutoff"], raw_rows=raw_rows, missing_customer_rows=missing_customers,
                 duplicate_occurrences_after_first=duplicates, raw_min_date=minimum.isoformat(), raw_max_date=maximum.isoformat(),
                 category_counts=dict(sorted(totals.items())), category_examples=dict(sorted(samples.items())),
                 coverage=dict(months=len(rows), country=config["country"], start=config["start"], end_exclusive=config["end_exclusive"]),
                 attribution="Chen, D. (2015). Online Retail. UCI Machine Learning Repository. DOI:10.24432/C5BW33. CC BY 4.0.",
                 boundaries=["Illustrative internal reporting policy, not UCI accounting guidance.",
                             "Amounts are recorded quantity times price, not recognized revenue, cash or profit.",
                             "Credits use posting month; no original-invoice linkage is supplied.",
                             "December 2011 is incomplete and excluded from full-month comparison.",
                             "Exact duplicate-looking records are retained because no unique source line identifier is supplied."])
    (output / "audit.json").write_text(json.dumps(audit, ensure_ascii=False, sort_keys=True, indent=2)+"\n")
    monthly_rows = [[r[0], *r[1:6], *[float(Decimal(v)) if v else None for v in r[6:]]] for r in rows]
    method = [
        ["批准版本", version, "V2仅内部分析；V3运营草案未授权，补签发生且收到于截点后。"],
        ["日期", f"{config['start']} <= InvoiceDate < {config['end_exclusive']}", "原始数据到2011-12-09，2011年12月不与完整月比较；不推断日历外未发生交易。"],
        ["交易净额", "sale_gbp-credit_gbp", "C前缀且负数量为本期贷项，按发票生成月入账；没有原单对应字段，不倒填原销售月。不是会计收入/利润/现金。"],
        ["缺失客户号", "金额保留；已知客户人数不计缺号", "官网无缺失标记与实际文件不一致，以原始文件实查；不得补客户号或删金额。"],
        ["重复外观", "保留全部原始行；仅报告重复出现数", "同InvoiceNo可有多商品；没有源行唯一键，不能把完全相同行擅自去重。"],
        ["符号及价格", "正价、非零数量；C与符号一致才纳入", "无C负数量及非正价等隔离，不据此断言都是客户退货；优先顺序见技能。"],
        ["当期贷项比", "credit_gbp/sale_gbp", "分子分母均为当期金额，非同批顾客退货率；不能据此推断获客/质量原因。"],
        ["平均正向发票额", "sale_gbp/正向独立InvoiceNo数", "不是除以商品行数，也不是把贷项净额除以发票数。"],
        ["来源与许可", audit["attribution"], "UCI原始工作簿按字节交付；本月报为原创演练变换，不代表UCI或数据捐赠者背书。"],
        ["复现", "python analysis.py --inputs /app/input_files --output /app/output --skills /skills", "先检查audit.json收支及类别守恒，再使用月报；仅内部复核未对外发布。"],
    ]
    quality = [[key, count] for key, count in sorted(totals.items())]
    quality += [["raw_rows", raw_rows], ["missing_customer_rows", missing_customers], ["duplicate_occurrences_after_first", duplicates]]
    money_columns = [0,6,7,8,9,11,12]
    volume_columns = [0,1,2,3,4,5,10]
    office_tools.workbook(output / "review.xlsx", [
        ("Monthly", [columns[i] for i in money_columns], [[r[i] for i in money_columns] for r in monthly_rows], [14,21,21,21,20,23,23]),
        ("Volume", [columns[i] for i in volume_columns], [[r[i] for i in volume_columns] for r in monthly_rows], [14,17,17,19,19,23,23]),
        ("Quality", ["类别或核对项", "行数"], quality, [60,25]),
        ("Method", ["事项", "采用口径", "适用范围与解释"], method, [19,48,70])], rows_per_page=None,
        compact_formula_rows=True, number_formats={"Monthly": {**{c:"#,##0.00" for c in ("B","C","D","F")}, **{c:"0.00%" for c in ("E","G")}}, "Volume":{"G":"#,##0.00"}})
    net = sum(Decimal(row[8]) for row in rows)
    peak = max(rows, key=lambda row: Decimal(row[8]))
    missing_value = sum(Decimal(row[11]) for row in rows)
    sales_value = sum(Decimal(row[6]) for row in rows)
    office_tools.document(output / "handoff.docx", "零售交易月报复核：内部交接", [
        ("可以使用的结果", [f"按{version}规则，UK的12个完整月交易净额合计{net:,.2f}英镑；净额最高为{peak[0]}，{Decimal(peak[8]):,.2f}英镑。该金额是记录中数量乘单价之和，不是确认收入、利润或已收款。",
                           f"保留缺失客户号的正向金额{missing_value:,.2f}英镑，占正向金额{missing_value/sales_value:.2%}；客户人数只统计已知ID。不能因为缺少ID就删去这部分交易金额。"]),
        ("旧稿需要纠正", ["不能先dropna再drop_duplicates并删掉所有取消单；这会同时改变金额、客户覆盖和贷项。多行同一发票也不能压成一行，发票数须独立计数。",
                        "C前缀表示取消，需结合负数量和正单价；非C负数量、零价和负价单列隔离。按批准顺序互斥计数，类别行数和必须等于原始行数。",
                        "原数据没有贷项对应的原销售发票号，客户号加商品号不足以证明唯一匹配；贷项放生成月份，不倒填原月。当期贷项金额比不是销售批次退货率。"]),
        ("时间与发布", ["资料截点2011-12-12 10:00 UTC。财务负责人V2已在截点前批准收到；运营V3无权更改口径，10:05补签不能回填。此稿仅供内部复核，2011年12月只到9日，不参与完整月比较；不能用月报确认2011年全年金额。"]),
        ("接手动作", ["数据管理员补取独立源行ID与贷项关联证据，才讨论去重和按原销售月追溯；财务负责人确认收入确认所需凭证，才编会计收入报表；运营负责人核对缺号流程和被隔离记录原因。当前月报继续交付，不把这些缺口当作全任务停摆。",
                     audit["attribution"], "复现：python analysis.py --inputs /app/input_files --output /app/output --skills /skills。使用固定本地文件；不联网刷新、补造数据、申报、联系客户或实际发布。"]),
    ])
    import shutil
    if Path(__file__).resolve() != (output / "analysis.py").resolve():
        shutil.copyfile(Path(__file__), output / "analysis.py")
    print(json.dumps(dict(raw_rows=raw_rows, categories=dict(totals), months=len(rows), net=str(net), duplicate_occurrences=duplicates)))


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--inputs", required=True, type=Path)
    parser.add_argument("--output", required=True, type=Path)
    parser.add_argument("--skills", required=True, type=Path)
    args = parser.parse_args()
    run(args.inputs, args.output, args.skills)
